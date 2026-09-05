#!/usr/bin/env python
"""전문가 HAZOP 워크시트(xlsx) → 골드셋 JSON 변환 CLI.

spec: `gold-dataset` T-01~T-11 · REQ-01~09

design.md 는 컴포넌트를 `tools/_gold/` 하위 7개 모듈로 나누지만, CLAUDE.md "압축 구현 원칙"
에 따라 이 파일 하나에 담는다. design.md 의 이름(XlsxLoader, ColumnMapper, GuidewordParser,
MultiValueParser, NodeParser, FieldValidator, DatasetValidator, NodeSplitter,
StratifiedSplitter, JsonWriter)은 그대로 유지해 추적성을 보존한다.

사용:
    PYTHONIOENCODING=utf-8 python tools/build_gold.py \\
        --input data/raw/D1_HAZOP_워크시트.xlsx --sheet HAZOP워크시트 \\
        --split node --tune-nodes N1
"""

from __future__ import annotations

import argparse
import json
import logging
import random
import re
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Final

import openpyxl
from jsonschema import Draft7Validator

if __package__ in (None, ""):  # `python tools/build_gold.py` 로 직접 실행한 경우
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tools._gold.models import (  # noqa: E402  (위 sys.path 보정 이후여야 한다)
    COLUMN_MAP,
    REQUIRED_COLUMNS,
    GoldRecord,
    NodeMeta,
    SplitConfig,
)

logger = logging.getLogger("build_gold")

REPO_ROOT: Final[Path] = Path(__file__).resolve().parent.parent
DEFAULT_INPUT: Final[str] = "data/raw/D1_HAZOP_워크시트.xlsx"
DEFAULT_SHEET: Final[str] = "HAZOP워크시트"
DEFAULT_RATING_SHEET: Final[str] = "평가기준"
DEFAULT_OUTPUT_DIR: Final[str] = "data/gold"
DEFAULT_PREFIX: Final[str] = "nh3"
SCHEMA_PATH: Final[Path] = REPO_ROOT / "schemas" / "gold_record.schema.json"


# ────────────────────────────────────────────────────────────────────────────
# XlsxLoader (REQ-01) — design.md §4.1
# ────────────────────────────────────────────────────────────────────────────
class XlsxLoader:
    """xlsx 시트를 헤더 기준 딕셔너리 리스트로 읽는다."""

    def load(self, path: Path, sheet: str) -> list[dict[str, Any]]:
        if not path.is_file():
            raise FileNotFoundError(f"input file not found: {path}")
        # data_only=True: 위험도 셀이 `=H*I` 수식이어도 캐시된 값으로 읽는다.
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"sheet '{sheet}' not found")
            rows_iter = workbook[sheet].iter_rows(values_only=True)
            header = next(rows_iter, None)
            if header is None:
                raise ValueError(f"sheet '{sheet}' is empty")
            keys = ["" if cell is None else str(cell).strip() for cell in header]
            rows: list[dict[str, Any]] = []
            for values in rows_iter:
                if _is_blank_row(values):
                    continue
                rows.append({k: v for k, v in zip(keys, values, strict=False) if k})
            logger.info("loaded %d rows from %s [%s]", len(rows), path, sheet)
            return rows
        finally:
            workbook.close()

    def load_sheet_matrix(self, path: Path, sheet: str) -> list[tuple[Any, ...]]:
        """헤더 해석 없이 시트를 행 튜플 그대로 읽는다(평가기준 시트용)."""
        if not path.is_file():
            raise FileNotFoundError(f"input file not found: {path}")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"sheet '{sheet}' not found")
            return [tuple(row) for row in workbook[sheet].iter_rows(values_only=True)]
        finally:
            workbook.close()


def _is_blank_row(values: tuple[Any, ...]) -> bool:
    return all(v is None or str(v).strip() == "" for v in values)


# ────────────────────────────────────────────────────────────────────────────
# ColumnMapper (REQ-02) — design.md §4.2
# ────────────────────────────────────────────────────────────────────────────
class ColumnMapper:
    """12열 헤더 존재를 확인하고 내부 필드명으로 rename 한다."""

    def validate_and_rename(self, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        present: set[str] = set()
        for row in rows:
            present.update(str(k).strip() for k in row)
        missing = [c for c in REQUIRED_COLUMNS if c not in present]
        if missing:
            logger.error("missing columns: %s", missing)
            raise ValueError(f"missing columns: {missing}")
        return [
            {COLUMN_MAP[str(k).strip()]: v for k, v in row.items() if str(k).strip() in COLUMN_MAP}
            for row in rows
        ]


# ────────────────────────────────────────────────────────────────────────────
# GuidewordParser (REQ-04) — design.md §4.3
# ────────────────────────────────────────────────────────────────────────────
_GW_PATTERN: Final[re.Pattern[str]] = re.compile(
    r"^(?P<gw>[^(（]+)\s*[(（](?P<param>[^)）]+)[)）]\s*$"
)

_GW_NORMALIZE: Final[dict[str, str]] = {
    "No": "No", "없음": "No", "없는": "No",
    "More": "More", "고": "More", "높은": "More", "증가": "More", "과다": "More",
    "Less": "Less", "저": "Less", "낮은": "Less", "감소": "Less", "부족": "Less",
    "Reverse": "Reverse", "역": "Reverse", "역방향": "Reverse", "반대": "Reverse",
    "Other than": "Other than", "이외": "Other than", "다른": "Other than",
    "Part of": "Part of", "일부": "Part of",
    "As well as": "As well as", "추가": "As well as", "이외에도": "As well as",
}


class GuidewordParser:
    """`"More (압력)"` → `("More", "압력")`. 미등록 가이드워드는 원본을 유지한다."""

    def parse(self, raw: str | None) -> tuple[str, str]:
        if raw is None or not str(raw).strip():
            return "", ""
        text = str(raw).strip()
        match = _GW_PATTERN.match(text)
        if match:
            guideword = match.group("gw").strip()
            parameter = match.group("param").strip()
        else:
            guideword, parameter = text, ""
        return _GW_NORMALIZE.get(guideword, guideword), parameter


# ────────────────────────────────────────────────────────────────────────────
# MultiValueParser (REQ-06) — design.md §4.4
# ────────────────────────────────────────────────────────────────────────────
_NUMBERED_LINE: Final[re.Pattern[str]] = re.compile(r"^\s*\d+[.)]\s*", re.MULTILINE)


class MultiValueParser:
    """원인·결과·안전장치·권고 셀을 구분자 우선순위(줄바꿈 → 번호 → 세미콜론)로 분해한다."""

    def parse(self, raw: str | None) -> list[str]:
        if raw is None or not str(raw).strip():
            return []
        text = str(raw)
        if "\n" in text:
            items = text.split("\n")
        elif _NUMBERED_LINE.search(text):
            items = _NUMBERED_LINE.split(text)
        elif ";" in text:
            items = text.split(";")
        else:
            items = [text]
        return [_NUMBERED_LINE.sub("", item).strip() for item in items if item.strip()]


# ────────────────────────────────────────────────────────────────────────────
# NodeParser (REQ-05) — design.md §4.5
# ────────────────────────────────────────────────────────────────────────────
_NODE_CODE: Final[re.Pattern[str]] = re.compile(r"^(N\d+)")
_SUBSTANCE_PATTERNS: Final[list[tuple[re.Pattern[str], str]]] = [
    (re.compile(r"NH\s*[₃3]|암모니아|ammonia", re.IGNORECASE), "NH3"),
    (re.compile(r"LNG|액화천연가스", re.IGNORECASE), "LNG"),
    (re.compile(r"LPG", re.IGNORECASE), "LPG"),
]
_PRESSURE: Final[re.Pattern[str]] = re.compile(
    r"(-?\d+(?:\.\d+)?)\s*(kPa|bar|MPa)\s*(?:g|G)?", re.IGNORECASE
)
_TEMPERATURE: Final[re.Pattern[str]] = re.compile(r"(-?\d+(?:\.\d+)?)\s*°?\s*C\b")
_EQUIPMENT_SPLIT: Final[re.Pattern[str]] = re.compile(r"[·,/]")
_PRESSURE_TO_KPA: Final[dict[str, float]] = {"kpa": 1.0, "bar": 100.0, "mpa": 1000.0}


class NodeParser:
    """노드 식별자 forward-fill 과 노드 메타 파싱(노드당 1회)."""

    def __init__(self, default_substance: str = "unknown") -> None:
        self._default_substance = default_substance
        self._current_node = "UNKNOWN"
        self._node_meta_cache: dict[str, NodeMeta] = {}

    def parse_node(self, raw: str | None) -> tuple[str, NodeMeta]:
        text = "" if raw is None else str(raw).strip()
        if text:
            self._current_node = self._extract_code(text)
            if self._current_node not in self._node_meta_cache:
                self._node_meta_cache[self._current_node] = self._parse_meta(text)
        if self._current_node not in self._node_meta_cache:
            # 첫 행부터 노드 셀이 비어 forward-fill 할 값이 없는 경우
            self._node_meta_cache[self._current_node] = NodeMeta(
                substance=self._default_substance
            )
        return self._current_node, self._node_meta_cache[self._current_node]

    def _extract_code(self, text: str) -> str:
        match = _NODE_CODE.match(text)
        if match:
            return match.group(1)
        logger.warning("node code not found in %r → UNKNOWN", text)
        return "UNKNOWN"

    def _parse_meta(self, text: str) -> NodeMeta:
        meta = NodeMeta(substance=self._default_substance)
        for pattern, name in _SUBSTANCE_PATTERNS:
            if pattern.search(text):
                meta.substance = name
                break
        meta.phase = _parse_phase(text)
        pressure = _PRESSURE.search(text)
        if pressure:
            unit = pressure.group(2).lower()
            meta.P_kPag = float(pressure.group(1)) * _PRESSURE_TO_KPA[unit]
        temperature = _TEMPERATURE.search(_normalize_minus(text))
        if temperature:
            meta.T_degC = float(temperature.group(1))
        remainder = _NODE_CODE.sub("", text).strip()
        meta.equipment = [p.strip() for p in _EQUIPMENT_SPLIT.split(remainder) if p.strip()]
        if meta.P_kPag is None and meta.T_degC is None:
            logger.debug("node meta: 운전조건(P/T) 파싱 불가 — %r", text)
        return meta


def _parse_phase(text: str) -> str:
    liquid = bool(re.search(r"액상|액화|liquid", text, re.IGNORECASE))
    gas = bool(re.search(r"기상|가스|증기|vapou?r|\bgas\b", text, re.IGNORECASE))
    if liquid and gas:
        return "liquid/gas"
    if liquid:
        return "liquid"
    if gas:
        return "gas"
    return "unknown"


def _normalize_minus(text: str) -> str:
    """유니코드 빼기표(U+2212)·전각 하이픈을 ASCII 하이픈으로 통일한다."""
    return text.replace("−", "-").replace("－", "-")


# ────────────────────────────────────────────────────────────────────────────
# FieldValidator (REQ-07) — design.md §4.6
# ────────────────────────────────────────────────────────────────────────────
def validate_sf(value: Any) -> int | None:
    """S·F 를 1–5 정수로 강제한다. 범위 밖·비정수는 None."""
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    if number != int(number):
        return None
    integer = int(number)
    return integer if 1 <= integer <= 5 else None


def compute_risk(s: int | None, f: int | None) -> int | None:
    return s * f if s is not None and f is not None else None


class FieldValidator:
    """행 단위 결측·이상값 처리. 건너뛸 행은 None 을 반환한다."""

    def __init__(self) -> None:
        self.null_sf_rows = 0
        self.risk_mismatch_rows = 0
        self.skipped_rows = 0
        self.empty_causes_rows = 0

    def validate_row(self, row: dict[str, Any]) -> dict[str, Any] | None:
        row_id = row.get("row_no")
        deviation = "" if row.get("deviation") is None else str(row["deviation"]).strip()
        guideword_raw = (
            "" if row.get("guideword_raw") is None else str(row["guideword_raw"]).strip()
        )
        if not deviation:
            logger.warning("row %s: deviation 이 비어 있어 건너뜀", row_id)
            self.skipped_rows += 1
            return None
        if not guideword_raw:
            logger.warning("row %s: 가이드워드가 비어 있어 건너뜀", row_id)
            self.skipped_rows += 1
            return None

        severity = validate_sf(row.get("severity"))
        frequency = validate_sf(row.get("frequency"))
        if severity is None or frequency is None:
            logger.warning(
                "row %s: S/F 범위 이상 (S=%r, F=%r) → null", row_id, row.get("severity"),
                row.get("frequency"),
            )
            self.null_sf_rows += 1

        risk = compute_risk(severity, frequency)
        raw_risk = validate_risk_cell(row.get("risk_score_raw"))
        if risk is not None and raw_risk is not None and raw_risk != risk:
            logger.warning("row %s: 위험도 셀 %s ≠ S×F %s → S×F 채택", row_id, raw_risk, risk)
            self.risk_mismatch_rows += 1

        scenario_cell = row.get("scenario")
        scenario = None if scenario_cell is None or not str(scenario_cell).strip() else str(
            scenario_cell
        ).strip()

        return {
            **row,
            "deviation": deviation,
            "guideword_raw": guideword_raw,
            "severity": severity,
            "frequency": frequency,
            "risk_score": risk,
            "scenario": scenario,
        }


def validate_risk_cell(value: Any) -> int | None:
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    return int(number) if number == int(number) else None


# ────────────────────────────────────────────────────────────────────────────
# RowTransformer — design.md §1 3단계
# ────────────────────────────────────────────────────────────────────────────
class RowTransformer:
    """매핑된 행 → `GoldRecord` 리스트."""

    def __init__(self, prefix: str = DEFAULT_PREFIX, substance: str = "unknown") -> None:
        self.prefix = prefix
        self.node_parser = NodeParser(default_substance=substance)
        self.guideword_parser = GuidewordParser()
        self.multi_value_parser = MultiValueParser()
        self.field_validator = FieldValidator()

    def transform(self, rows: list[dict[str, Any]]) -> list[GoldRecord]:
        records: list[GoldRecord] = []
        for index, row in enumerate(rows, start=1):
            node, node_meta = self.node_parser.parse_node(row.get("node_raw"))
            checked = self.field_validator.validate_row(row)
            if checked is None:
                continue
            guideword, parameter = self.guideword_parser.parse(checked["guideword_raw"])
            causes = self.multi_value_parser.parse(checked.get("causes_raw"))
            if not causes:
                logger.warning("row %s: causes 가 비어 있음", checked.get("row_no"))
                self.field_validator.empty_causes_rows += 1
            records.append(
                GoldRecord(
                    id=f"{self.prefix}-{_row_number(checked.get('row_no'), index):03d}",
                    node=node,
                    node_meta=node_meta,
                    guideword=guideword,
                    parameter=parameter,
                    deviation=checked["deviation"],
                    causes=causes,
                    consequences=self.multi_value_parser.parse(checked.get("consequences_raw")),
                    safeguards_before=self.multi_value_parser.parse(
                        checked.get("safeguards_raw")
                    ),
                    S=checked["severity"],
                    F=checked["frequency"],
                    risk_score=checked["risk_score"],
                    recommendations=self.multi_value_parser.parse(
                        checked.get("recommendations_raw")
                    ),
                    scenario=checked["scenario"],
                )
            )
        return records


def _row_number(raw: Any, fallback: int) -> int:
    """`No` 열을 id 번호로 쓴다. 정수로 읽히지 않으면 행 순번을 쓴다."""
    try:
        return int(float(raw))  # type: ignore[arg-type]
    except (TypeError, ValueError):
        logger.warning("row No %r 를 정수로 읽을 수 없어 순번 %d 사용", raw, fallback)
        return fallback


# ────────────────────────────────────────────────────────────────────────────
# DatasetValidator (REQ-03) — design.md §1 4단계
# ────────────────────────────────────────────────────────────────────────────
class DatasetValidator:
    """레코드 배열 전체를 JSON Schema 로 검증하고 id 유일성을 확인한다."""

    def validate(self, records: list[dict[str, Any]], schema_path: Path) -> list[str]:
        schema = json.loads(schema_path.read_text(encoding="utf-8"))
        validator = Draft7Validator(schema)
        errors: list[str] = []
        seen: set[str] = set()
        for record in records:
            record_id = str(record.get("id"))
            for error in sorted(validator.iter_errors(record), key=str):
                errors.append(f"{record_id}: {error.message}")
            if record_id in seen:
                errors.append(f"{record_id}: duplicate id")
            seen.add(record_id)
        if not records:
            errors.append("no records produced")
        return errors


# ────────────────────────────────────────────────────────────────────────────
# Splitter (REQ-08) — design.md §4.7·§4.8·§6
# ────────────────────────────────────────────────────────────────────────────
class NodeSplitter:
    def split(
        self, records: list[GoldRecord], tune_nodes: list[str]
    ) -> tuple[list[GoldRecord], list[GoldRecord]]:
        tune = [r for r in records if r.node in tune_nodes]
        holdout = [r for r in records if r.node not in tune_nodes]
        return tune, holdout


class StratifiedSplitter:
    """guideword 층화 70/30 분할.

    scikit-learn 은 쓰지 않는다(CLAUDE.md 범위 규율). 층별로 `random.Random(seed)` 셔플 후
    비율 분할하며, 같은 seed 는 같은 결과를 준다.
    """

    def split(
        self, records: list[GoldRecord], seed: int = 42, ratio: float = 0.7
    ) -> tuple[list[GoldRecord], list[GoldRecord]]:
        groups: dict[str, list[int]] = {}
        for index, record in enumerate(records):
            groups.setdefault(record.guideword, []).append(index)
        rng = random.Random(seed)
        tune_idx: set[int] = set()
        for key in sorted(groups):
            indices = list(groups[key])
            rng.shuffle(indices)
            take = max(1, round(len(indices) * ratio)) if indices else 0
            tune_idx.update(indices[:take])
        tune = [r for i, r in enumerate(records) if i in tune_idx]
        holdout = [r for i, r in enumerate(records) if i not in tune_idx]
        return tune, holdout


def decide_splitter(
    records: list[GoldRecord],
    split_mode: str,
    tune_nodes: list[str] | None,
    seed: int,
) -> tuple[list[GoldRecord], list[GoldRecord], SplitConfig]:
    """design.md §6 결정 트리. 노드 수 < 2 이면 stratified 로 자동 전환한다."""
    nodes = sorted({r.node for r in records})
    timestamp = datetime.now(UTC).isoformat()
    if split_mode == "node" and len(nodes) >= 2:
        selected = tune_nodes or nodes[:1]
        tune, holdout = NodeSplitter().split(records, selected)
        config = SplitConfig(
            split_type="node",
            tune_nodes=selected,
            seed=None,
            tune_count=len(tune),
            eval_count=len(holdout),
            timestamp=timestamp,
        )
        return tune, holdout, config

    fallback_reason = None
    if split_mode == "node":
        fallback_reason = f"고유 노드 수 {len(nodes)} < 2 → stratified 자동 전환"
        logger.warning("%s", fallback_reason)
    tune, holdout = StratifiedSplitter().split(records, seed=seed)
    config = SplitConfig(
        split_type="stratified",
        tune_nodes=None,
        seed=seed,
        tune_count=len(tune),
        eval_count=len(holdout),
        timestamp=timestamp,
        fallback_reason=fallback_reason,
    )
    return tune, holdout, config


# ────────────────────────────────────────────────────────────────────────────
# JsonWriter (REQ-03·08·09) — design.md §4 6단계
# ────────────────────────────────────────────────────────────────────────────
class JsonWriter:
    """UTF-8(BOM 없음), 들여쓰기 2칸, `ensure_ascii=False` 로 직렬화한다."""

    def write(self, payload: Any, path: Path, dry_run: bool = False) -> None:
        if dry_run:
            logger.info("[dry-run] %s 미생성", path)
            return
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        logger.info("wrote %s", path)


def write_split_config(
    writer: JsonWriter, config: SplitConfig, path: Path, dry_run: bool = False
) -> None:
    writer.write(config.to_dict(), path, dry_run=dry_run)


def build_split_node(
    config: SplitConfig, tune: list[GoldRecord], holdout: list[GoldRecord]
) -> dict[str, Any]:
    """지시문 요구 산출물 `split_node.json` — 노드 단위 홀드아웃 명세(레코드 id 목록)."""
    return {
        "split_type": config.split_type,
        "tune_nodes": config.tune_nodes,
        "holdout_nodes": sorted({r.node for r in holdout}),
        "seed": config.seed,
        "tune_count": len(tune),
        "holdout_count": len(holdout),
        "timestamp": config.timestamp,
        "fallback_reason": config.fallback_reason,
        "tune_ids": [r.id for r in tune],
        "holdout_ids": [r.id for r in holdout],
    }


# ────────────────────────────────────────────────────────────────────────────
# 평가기준 시트 → rating_scale.json
# ────────────────────────────────────────────────────────────────────────────
_RISK_BAND: Final[re.Pattern[str]] = re.compile(r"^\d+\s*~\s*\d+$")


def parse_rating_scale(matrix: list[tuple[Any, ...]], sheet: str) -> dict[str, Any]:
    """`평가기준` 시트를 S/F 등급 정의와 위험도 구간으로 정리한다(원문 그대로 보존)."""
    severity: list[dict[str, Any]] = []
    frequency: list[dict[str, Any]] = []
    risk_bands: list[dict[str, str]] = []
    notes: list[str] = []
    for row in matrix:
        cells = ["" if c is None else str(c).strip() for c in row]
        cells += [""] * (3 - len(cells)) if len(cells) < 3 else []
        first, second, third = cells[0], cells[1], cells[2]
        if not first:
            continue
        if first.startswith("S "):
            severity.append({"grade": _to_int(second), "definition": third})
        elif first.startswith("F "):
            frequency.append({"grade": _to_int(second), "definition": third})
        elif _RISK_BAND.match(first):
            risk_bands.append({"range": first, "judgement": second, "action": third})
        elif first.startswith("주"):
            notes.append(first)
    return {
        "source_sheet": sheet,
        "severity": severity,
        "frequency": frequency,
        "risk_bands": risk_bands,
        "notes": notes,
    }


def _to_int(value: str) -> int | None:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


# ────────────────────────────────────────────────────────────────────────────
# CLI (REQ-01·09)
# ────────────────────────────────────────────────────────────────────────────
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="build_gold",
        description="전문가 HAZOP 워크시트(xlsx) → 골드셋 JSON (spec:gold-dataset)",
    )
    parser.add_argument(
        "--input", required=True, type=Path, help=f"원본 xlsx 경로 (권장: {DEFAULT_INPUT})"
    )
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"워크시트 시트명 (기본 {DEFAULT_SHEET})")
    parser.add_argument(
        "--rating-sheet",
        default=DEFAULT_RATING_SHEET,
        help=f"S/F 등급 정의 시트명 (기본 {DEFAULT_RATING_SHEET}, 없으면 건너뜀)",
    )
    parser.add_argument("--output-dir", default=Path(DEFAULT_OUTPUT_DIR), type=Path)
    parser.add_argument("--split", choices=["node", "stratified"], default="node")
    parser.add_argument("--tune-nodes", default=None, help="예: N1 또는 N1,N2")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--prefix", default=DEFAULT_PREFIX, help="레코드 id·출력 파일 접두어")
    parser.add_argument("--substance", default="NH3", help="노드 메타 substance 기본값")
    parser.add_argument("--dry-run", action="store_true", help="파일을 쓰지 않고 요약만 출력")
    parser.add_argument("--verbose", action="store_true", help="DEBUG 로깅")
    return parser


def run(args: argparse.Namespace) -> int:
    input_path: Path = args.input
    if "data/raw" not in input_path.as_posix():  # NFR-G05
        logger.warning("원본 경로가 data/raw/ 밖입니다: %s", input_path)

    loader = XlsxLoader()
    rows = ColumnMapper().validate_and_rename(loader.load(input_path, args.sheet))

    transformer = RowTransformer(prefix=args.prefix, substance=args.substance)
    records = transformer.transform(rows)
    payload = [r.to_dict() for r in records]

    errors = DatasetValidator().validate(payload, SCHEMA_PATH)
    if errors:
        for message in errors[:20]:
            logger.error("schema: %s", message)
        raise ValueError(f"schema validation failed: {len(errors)} error(s)")

    tune_nodes = [n.strip() for n in str(args.tune_nodes).split(",") if n.strip()] if args.tune_nodes else None
    tune, holdout, config = decide_splitter(records, args.split, tune_nodes, args.seed)

    out_dir: Path = args.output_dir
    stem = f"hazop_{args.prefix}"
    writer = JsonWriter()
    writer.write(payload, out_dir / f"{stem}.json", dry_run=args.dry_run)
    writer.write([r.to_dict() for r in tune], out_dir / f"{stem}_tune.json", dry_run=args.dry_run)
    writer.write([r.to_dict() for r in holdout], out_dir / f"{stem}_eval.json", dry_run=args.dry_run)
    writer.write(build_split_node(config, tune, holdout), out_dir / "split_node.json", dry_run=args.dry_run)
    write_split_config(writer, config, out_dir / "split_config.json", dry_run=args.dry_run)

    rating_written = False
    if args.rating_sheet:
        try:
            matrix = loader.load_sheet_matrix(input_path, args.rating_sheet)
        except ValueError:
            logger.warning("등급 정의 시트 '%s' 없음 — rating_scale.json 생략", args.rating_sheet)
        else:
            writer.write(
                parse_rating_scale(matrix, args.rating_sheet),
                out_dir / "rating_scale.json",
                dry_run=args.dry_run,
            )
            rating_written = True

    _print_summary(records, tune, holdout, transformer, out_dir, stem, rating_written)
    return 0


def _print_summary(
    records: list[GoldRecord],
    tune: list[GoldRecord],
    holdout: list[GoldRecord],
    transformer: RowTransformer,
    out_dir: Path,
    stem: str,
    rating_written: bool,
) -> None:
    validator = transformer.field_validator
    nodes = ", ".join(sorted({r.node for r in records}))
    null_sf = sum(1 for r in records if r.S is None or r.F is None)
    print(
        f"[build_gold] 완료: {len(records)} 레코드 | 노드: {nodes} | "
        f"null S/F: {null_sf} | 건너뜀: {validator.skipped_rows}"
    )
    print(f"  → {out_dir / f'{stem}.json'}")
    print(f"  → {out_dir / f'{stem}_tune.json'} ({_nodes_of(tune)}, {len(tune)} 레코드)")
    print(f"  → {out_dir / f'{stem}_eval.json'} ({_nodes_of(holdout)}, {len(holdout)} 레코드)")
    print(f"  → {out_dir / 'split_node.json'}")
    print(f"  → {out_dir / 'split_config.json'}")
    if rating_written:
        print(f"  → {out_dir / 'rating_scale.json'}")
    if validator.risk_mismatch_rows:
        print(f"  주의: 위험도 셀 불일치 {validator.risk_mismatch_rows}건 → S×F 로 재계산")
    if validator.empty_causes_rows:
        print(f"  주의: causes 빈 레코드 {validator.empty_causes_rows}건")


def _nodes_of(records: list[GoldRecord]) -> str:
    return ", ".join(sorted({r.node for r in records})) or "-"


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s | %(levelname)-8s | %(name)s | %(message)s",
    )
    # CLAUDE.md 규칙 8: 한글 출력은 UTF-8 가정. cp949 콘솔에서도 죽지 않도록 보정한다.
    if (sys.stdout.encoding or "").lower() not in ("utf-8", "utf8"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]
    try:
        return run(args)
    except (FileNotFoundError, ValueError) as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
