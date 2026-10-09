"""추론 경계 Z-1~Z-3 오프라인 시험 (PRD_본선_추론경계_Kiro연동 §3).

Z-1 정보 4단계 · Z-2 노드 유형별 경계표(`data/kb/inference_policy.json`) · Z-3 셀 보류("정보 부족").
전부 `MockBedrockClient` 경로 — 실호출 없음.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from jsonschema import Draft7Validator
from openpyxl import load_workbook
from pydantic import ValidationError

from apps.web import guide, service
from apps.web.replay import Result
from core import policy
from core.agent import DeviationRecord, GeneratorConfig, HazopGenerator, NodeMeta, verify
from core.agent.generate import INSUFFICIENT, _load_prompt, _split_prompt
from core.criteria import OFFICIAL_CRITERIA
from core.export import build_report, export_all, normalize_rows, render_lopa
from core.llm import ConverseResponse, Message, MockBedrockClient
from core.llm.config import _parse

_ROOT = Path(__file__).parent.parent
_SCHEMA = json.loads((_ROOT / "schemas" / "deviation.schema.json").read_text(encoding="utf-8"))
_PARAMS = ["유량", "압력", "온도", "준위", "조성", "부식"]
_META = NodeMeta(node="X1", substance="프로판", phase="liquid/gas", equipment=["LPG 저장탱크", "이송 펌프"])


def _factory(held: set[str], extra: dict[str, Any] | None = None) -> Any:
    """`held` 파라미터 셀은 보류(insufficient)로, 나머지는 판정으로 돌려준다. 보류 셀에도 원인·S·F 를 일부러 넣는다."""

    def make(system: str, messages: list[Message], **_: Any) -> ConverseResponse:
        if "파라미터 축" in system:
            return ConverseResponse(content=json.dumps({"parameters": [{"name": n, "rationale": "r"} for n in _PARAMS]}))
        user = str(messages[0].content)
        gw = next(g for g in ("No", "More", "Less", "Reverse", "Other than", "Part of", "As well as") if f"\n{g} —" in user)
        cells = []
        for n in _PARAMS:
            cell: dict[str, Any] = {"parameter": n, "applicable": True, "deviation": f"{n} 이탈", "causes": ["밸브 오조작"],
                                    "consequences": ["누출"], "safeguards_before": [], "S": 3, "F": 2,
                                    "recommendations": ["유무 확인"], "evidence": [], "confidence": "inferred"}
            if n in held:
                cell.update({"insufficient": True, "missing": ["액위 경보 설정값", "방유제 유무"], **(extra or {})})
            cells.append(cell)
        return ConverseResponse(content=json.dumps({"guideword": gw, "cells": cells}, ensure_ascii=False))

    return make


def _run(boundary: bool, held: set[str] = frozenset(), meta: NodeMeta = _META, **kw: Any) -> tuple[HazopGenerator, MockBedrockClient, list[DeviationRecord]]:  # noqa: B006
    client = MockBedrockClient(response_factory=_factory(set(held), kw.get("extra")))
    gen = HazopGenerator(client, GeneratorConfig(inference_boundary=boundary), criteria_id=OFFICIAL_CRITERIA)
    return gen, client, gen.generate(meta)


# ── Z-1·Z-2 정책 파일 → 프롬프트 ──────────────────────────────────────────────
def test_node_types_by_keyword_plus_common() -> None:
    types = [t["id"] for t in policy.node_types("X1", ["LPG 저장탱크", "이송 펌프"])]
    assert types == ["storage", "rotating", "common"]
    assert [t["id"] for t in policy.node_types("X1", ["무언가"])] == ["common"]


def test_boundary_off_keeps_prompts_and_schema() -> None:
    _, client, _ = _run(False)
    enum, judge = client.calls[0], client.calls[1]
    assert "부족하다고 적지도 마라" in enum["system"]
    assert "정보 단계" not in judge["system"] and "추론 경계" not in judge["messages"][0].content
    assert "insufficient" not in judge["response_schema"]["properties"]["cells"]["items"]["properties"]
    assert "- 기존 안전장치: 없음" in judge["messages"][0].content  # Z 이전 그대로


def test_boundary_on_puts_policy_table_into_prompts() -> None:
    """Z-G2: 표의 금지(U) 항목이 판정 프롬프트에, 4단계 정의가 시스템에. 열거의 '부족하다고 적지도 마라'는 빠진다."""
    _, client, _ = _run(True)
    enum, judge = client.calls[0], client.calls[1]
    assert "부족하다고 적지도 마라" not in enum["system"] and "'정보 부족'으로 보류" in enum["system"]
    assert policy.levels_text() in judge["system"]
    user = judge["messages"][0].content
    for t in policy.node_types(_META.node, _META.equipment):
        assert all(u in user for u in t["unknown"]) and all(i in user for i in t["infer"])
    assert "반응열" not in user  # 걸리지 않은 유형(반응기)은 넣지 않는다
    assert "'없음'이 아니라 '모름'" in user
    props = judge["response_schema"]["properties"]["cells"]["items"]["properties"]
    assert props["insufficient"] == {"type": "boolean"} and props["missing"]["maxItems"] == 3


def test_boundary_system_block_is_node_independent() -> None:
    """시스템 블록은 노드와 무관해야 캐시가 산다(R-08) — 노드별 경계표는 사용자 턴에만."""
    _, a, _ = _run(True)
    _, b, _ = _run(True, meta=NodeMeta(node="Y9", substance="수소", phase="gas", equipment=["반응기"]))
    assert a.calls[1]["system"] == b.calls[1]["system"]


@pytest.mark.parametrize(("safeguards", "known", "text"), [
    ([], False, "미상 (입력에 없음 — '없음'이 아니라 '모름')"),
    ([], True, "없음 (입력에 명시)"),
    (["안전밸브"], False, "안전밸브"),
])
def test_safeguards_none_vs_unknown(safeguards: list[str], known: bool, text: str) -> None:
    assert policy.safeguards_text(safeguards, known) == text


def test_enumerate_rule_must_match_prompt_file() -> None:
    """프롬프트 파일의 줄이 바뀌어 치환이 조용히 안 되는 일이 없게 — 못 찾으면 예외."""
    system = _split_prompt(_load_prompt("matrix_enumerate.md"))[0]
    assert policy.enumerate_rule(system) != system
    with pytest.raises(ValueError, match="바꿀 줄"):
        policy.enumerate_rule(system.replace("부족하다고 적지도 마라", "부족하다고 적지 마라"))


# ── Z-3 셀 보류 ───────────────────────────────────────────────────────────────
def test_held_cell_drops_causes_and_grades() -> None:
    gen, _, records = _run(True, held={"준위"})
    held = [r for r in records if r.status == INSUFFICIENT]
    assert len(held) == 7  # 가이드워드 7종 × '준위'
    r = held[0]
    assert (r.S, r.F, r.risk_score, r.causes) == (None, None, None, [])
    assert r.missing == ["액위 경보 설정값", "방유제 유무"] and r.deviation == "준위 이탈"
    assert gen.judged_cells == gen.expected_cells == 42  # 보류도 판정한 셀로 센다
    assert all(x.S is not None for x in records if x.status != INSUFFICIENT)


def test_insufficient_ignored_when_boundary_off() -> None:
    """끈 실행에서는 스키마에 없는 키다 — 모델이 넣어도 보류하지 않는다(판정 그대로)."""
    _, _, records = _run(False, held={"준위"})
    assert all(r.status == "applicable" and r.S == 3 for r in records)


def test_record_without_grades_must_be_held() -> None:
    with pytest.raises(ValidationError, match="보류"):
        DeviationRecord(id="x1-001", node="X1", node_meta=_META, guideword="No", parameter="p", deviation="d")


def test_held_record_matches_output_schema() -> None:
    _, _, records = _run(True, held={"준위"})
    out, _ = verify(records)
    errors = list(Draft7Validator(_SCHEMA).iter_errors([r.model_dump() for r in out]))
    assert not errors, errors[:3]


def test_verify_tiers_held_rows() -> None:
    _, _, records = _run(True, held={"준위"}, extra={"consequences": ["반경 30 m 피해"]})
    out, summary = verify(records)
    held = [r for r in out if r.status == INSUFFICIENT]
    assert {r.confidence for r in held} == {"review"}  # 근거 없는 수치는 보류 행이어도 🔴
    _, _, clean = _run(True, held={"준위"})
    assert {r.confidence for r in verify(clean)[0] if r.status == INSUFFICIENT} == {INSUFFICIENT}
    assert summary.by_rule["unsupported_number"] == 7


# ── 내보내기 ──────────────────────────────────────────────────────────────────
def test_export_blanks_risk_and_adds_held_sheet(tmp_path: Path) -> None:
    _, _, records = _run(True, held={"준위"})
    out, _ = verify(records)
    paths = export_all(out, tmp_path)
    wb = load_workbook(paths["xlsx"])
    ws = wb["HAZOP워크시트"]
    rows = list(ws.iter_rows(min_row=2, max_row=len(out) + 1, values_only=True))
    held_rows = [r for r in rows if r[7] is None]
    assert len(held_rows) == 7 and all(r[8] is None and r[9] is None for r in held_rows)
    assert all(str(r[9]).startswith("=INDEX") for r in rows if r[7] is not None)
    sheet = wb["확인 필요"]
    assert [c.value for c in sheet[1]] == ["No", "가이드워드", "이탈 초안", "필요한 정보"]
    assert sheet.max_row == 8 and sheet[2][3].value == "액위 경보 설정값 · 방유제 유무"
    assert "정보 부족" in str(ws.cell(row=len(out) + 3, column=1).value)
    report = build_report(normalize_rows([r.model_dump() for r in out]))
    assert report.insufficient_count == 7 and report.confidence_distribution["insufficient"] == 7
    assert sum(sum(v.values()) for v in report.sf_matrix.values()) == len(out) - 7
    assert "준위 이탈" not in render_lopa(normalize_rows([r.model_dump() for r in out]), top_n=100)


def test_export_without_held_rows_has_no_extra_sheet(tmp_path: Path) -> None:
    _, _, records = _run(True)
    wb = load_workbook(export_all(verify(records)[0], tmp_path)["xlsx"])
    assert "확인 필요" not in wb.sheetnames
    assert build_report(normalize_rows([r.model_dump() for r in records])).insufficient_count is None


# ── 화면 ──────────────────────────────────────────────────────────────────────
def _result(held: set[str] = frozenset(), safeguards_known: bool = False) -> Result:  # noqa: B006
    meta_node = _META.model_copy(update={"safeguards_known": safeguards_known})
    gen, _, records = _run(True, held=set(held), meta=meta_node)
    meta = {"source": "mock", "node": "X1", "split": "none", "node_meta": meta_node.model_dump(),
            "criteria_id": OFFICIAL_CRITERIA, "expected_cells": gen.expected_cells, "judged_cells": gen.judged_cells,
            "inference_boundary": gen.inference_boundary}
    return Result(meta=meta, records=records)


def test_screen_labels_held_rows_and_counts() -> None:
    result = _result({"준위"})
    table = service.worksheet_table(result)
    held = [row for row in table if row["S(1-5)"] is None]
    assert len(held) == 7 and held[0][service.CONFIDENCE_COLUMN] == "⚪ 정보 부족 — 액위 경보 설정값 · 방유제 유무 미상"
    assert held[0]["위험도"] is None
    assert service.held_rows(result)[0]["필요한 정보"] == "액위 경보 설정값 · 방유제 유무"
    assert service.cell_breakdown(result) == "판정 42셀 = 해당 35 · 해당 없음 0 · 정보 부족 7(보류)"
    assert "정보 부족 7(보류)" in service.summary_line(result)
    assert service.confidence_counts(result).endswith("⚪ 정보 부족 7")
    assert service.f_distribution(result) == "F=2 비율 100%"


def test_safeguards_notice_only_when_unknown() -> None:
    assert service.safeguards_notice(_result()) == service.SAFEGUARDS_UNKNOWN_NOTICE
    assert service.safeguards_notice(_result(safeguards_known=True)) is None
    off = _result()
    off.meta["inference_boundary"] = False
    assert service.safeguards_notice(off) is None and service.cell_breakdown(off) is None


def test_review_can_fill_held_grades_but_not_blank_judged() -> None:
    result = _result({"준위"})
    table = service.worksheet_table(result)
    i_held = next(i for i, row in enumerate(table) if row["S(1-5)"] is None)
    kept, log = service.apply_review(result, {i_held: {"S(1-5)": 3, "F(1-5)": 2}})
    assert (kept[i_held]["S"], kept[i_held]["F"]) == (3, 2) and log[0][1] == "수정"
    i_ok = next(i for i, row in enumerate(table) if row["S(1-5)"] is not None)
    with pytest.raises(ValueError, match="비울 수 없"):
        service.apply_review(result, {i_ok: {"S(1-5)": None}})
    assert service.apply_review(result, {i_held: {"S(1-5)": float("nan")}})[1] == []  # 빈칸 그대로는 수정 아님


def test_guide_levels_use_policy_words() -> None:
    """Z-G1: 화면 4단계 표가 프롬프트와 같은 문구(정책 파일)."""
    rows = guide.information_level_rows()
    assert [r["단계"] for r in rows] == ["G 주어짐", "S 근거 있음", "I 일반 추론", "U 미상"]
    assert all(r["뜻"] in policy.levels_text() for r in rows)


# ── 설정 ──────────────────────────────────────────────────────────────────────
def test_config_flag_read_and_validated() -> None:
    raw = {"provider": "anthropic", "generation": {"model_id": "m", "temperature": 0.2, "max_tokens": 10,
                                                   "inference_boundary": True},
           "verifier": {"model_id": "v", "temperature": 0.0, "max_tokens": 10}}
    assert _parse(raw).inference_boundary is True
    raw["generation"]["inference_boundary"] = "false"
    from core.llm import ConfigValidationError

    with pytest.raises(ConfigValidationError):
        _parse(raw)
