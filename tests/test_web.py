"""오프라인 시험 — FR-10 H-04 (지시문 H (a)~(g)) + H-06 (지시문 I-1 (a)~(d)) + J-01·J-03·J-04 (지시문 J). 네트워크 0회."""

from __future__ import annotations

import io
import json
from datetime import date
from typing import TYPE_CHECKING

import openpyxl
import pytest
from jsonschema import Draft7Validator

from apps.web import service
from apps.web.replay import REPLAY_DIR, load_replay, load_replays
from core.agent.generate import STANDARD_GUIDEWORDS

if TYPE_CHECKING:
    from pathlib import Path

_ROOT = REPLAY_DIR.parents[1]
_SCHEMA = json.loads((_ROOT / "schemas" / "deviation.schema.json").read_text(encoding="utf-8"))
_APP = _ROOT / "apps" / "web" / "app.py"


@pytest.fixture(autouse=True)
def _reset_daily_counter() -> None:
    service._daily_runs.clear()


def _write(path: Path, source: str, captured_at: str) -> None:
    # 첫 파일(glob 순서)이 아니라 N1 파일을 명시 — 10/8 i1_*.json 이 n1_ 보다 앞에 정렬돼 노드가 I1 로 바뀌었다.
    payload = json.loads((REPLAY_DIR / "n1_20260929.json").read_text(encoding="utf-8"))
    payload.update(source=source, captured_at=captured_at)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


# (a) live 가 gold 보다 우선 — 파일명·시각 순서가 gold 쪽에 유리해도 live 를 고른다.
def test_load_replay_prefers_live_over_gold(tmp_path: Path) -> None:
    _write(tmp_path / "a_live.json", "live", "2026-09-01T00:00:00+00:00")
    _write(tmp_path / "z_gold.json", "gold", "2026-09-28T00:00:00+00:00")
    assert load_replay(tmp_path).meta["source"] == "live"
    (tmp_path / "a_live.json").unlink()
    assert load_replay(tmp_path).meta["source"] == "gold"


# (b) 커밋된 재생 파일이 스키마를 100% 통과하고 risk_score == S×F.
def test_replay_records_are_schema_valid() -> None:
    result = load_replay()
    assert result.records
    payload = [r.model_dump() for r in result.records]
    assert not list(Draft7Validator(_SCHEMA).iter_errors(payload))
    assert all(r.risk_score == r.S * r.F for r in result.records)


# (c) 재생 → export_all 3개 파일, 워크시트 데이터 행 수 == records 수.
def test_export_files_from_replay() -> None:
    result = load_replay()
    files = service.export_files(result)
    assert set(files) == {"xlsx", "lopa", "report"}
    assert all(data for _, data in files.values())
    ws = openpyxl.load_workbook(io.BytesIO(files["xlsx"][1]))["HAZOP워크시트"]
    data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]
    assert len(data_rows) == len(result.records)
    assert len(service.worksheet_table(result)) == len(result.records)


# (d) 실호출 활성 판정.
@pytest.mark.parametrize(
    ("environ", "enabled"),
    [
        ({"ANTHROPIC_API_KEY": "k"}, False),  # ALLOW 없음
        ({"HAZOP_ALLOW_LIVE": "true"}, False),  # 키 없음
        ({"HAZOP_ALLOW_LIVE": "true", "ANTHROPIC_API_KEY": "  "}, False),  # 빈 키
        ({"HAZOP_ALLOW_LIVE": "false", "ANTHROPIC_API_KEY": "k"}, False),
        ({"HAZOP_ALLOW_LIVE": "true", "ANTHROPIC_API_KEY": "k"}, True),
    ],
)
def test_live_mode_gate(environ: dict[str, str], enabled: bool) -> None:
    assert (service.live_block_reason(environ) is None) is enabled


def test_sync_secrets_copies_into_environ() -> None:
    environ: dict[str, str] = {}
    service.sync_secrets({"HAZOP_ALLOW_LIVE": "true", "ANTHROPIC_API_KEY": "k", "X": "y"}, environ)
    assert environ == {"HAZOP_ALLOW_LIVE": "true", "ANTHROPIC_API_KEY": "k"}
    assert service.live_block_reason(environ) is None


# (e) 상한 — 세션 1회·일 5회.
def test_quota_rejects_second_session_run_and_sixth_daily_run() -> None:
    day = date(2026, 9, 29)
    assert service.reserve_live_run(0, day) is None
    reason = service.reserve_live_run(1, day)
    assert reason and "세션" in reason
    for _ in range(service.DAILY_LIMIT - 1):
        assert service.reserve_live_run(0, day) is None
    reason = service.reserve_live_run(0, day)
    assert reason and "5회" in reason
    assert service.reserve_live_run(0, date(2026, 9, 30)) is None  # 날짜가 바뀌면 초기화


# (f) mock + ALLOW 로 실호출 경로가 끝까지 돈다.
def test_live_path_runs_end_to_end_on_mock() -> None:
    environ = {"HAZOP_USE_MOCK": "true", "HAZOP_ALLOW_LIVE": "true"}
    assert service.live_block_reason(environ) is None
    replay = load_replay()
    result = service.run_live(json.dumps(replay.meta["node_meta"]), replay, environ)
    assert result.meta["source"] == "mock"
    assert result.meta["judged_cells"] == result.meta["expected_cells"] > 0
    assert result.meta["review_guidewords"] == []
    assert {(r.guideword, r.parameter) for r in result.records} == {
        (r.guideword, r.parameter) for r in replay.records if r.guideword in STANDARD_GUIDEWORDS
    }
    assert "판정 셀" in service.summary_line(result)
    assert set(service.export_files(result)) == {"xlsx", "lopa", "report"}


# (g) AppTest — 사례 모드로 바꾸면 첫 공정의 N1 결과표가 바로 보이고, 노드 버튼으로 바뀐다. 예외 0건.
# (J-04 에서 "프리셋을 눌러야 표가 뜬다" → "공정을 고르면 첫 노드 표가 뜬다" 로 바뀌었다.)
def test_app_preset_click_shows_table(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    monkeypatch.delenv("HAZOP_ALLOW_LIVE", raising=False)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    at.radio(key="mode").set_value("실측 사례 재생").run()
    assert len(at.dataframe) == 1
    assert len(at.dataframe[0].value) == len(load_replay().records)
    at.button(key="preset_N2").click().run()
    assert not at.exception
    assert len(at.dataframe[0].value) == len(load_replays()["N2"].records)


# ── H-06 (지시문 I-1) ─────────────────────────────────────────────────────────
def _write_node(
    path: Path,
    node: str,
    source: str,
    captured_at: str,
    *,
    recall: tuple[int, int] | None = None,
    recall_key: str = "recall",
) -> None:
    payload = json.loads((REPLAY_DIR / "n1_20260929.json").read_text(encoding="utf-8"))
    payload.pop("recall_n1", None)
    payload.update(
        node=node,
        source=source,
        captured_at=captured_at,
        split="tune" if node == "N1" else "holdout",
        records=payload["records"][:3],
    )
    payload[recall_key] = (
        None if recall is None else {"recall": recall[0] / recall[1], "matched": recall[0], "total": recall[1]}
    )
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


# (a) 노드별 우선순위 — 노드마다 따로 live > gold, 같은 source 는 captured_at 최신.
def test_load_replays_priority_is_per_node(tmp_path: Path) -> None:
    _write_node(tmp_path / "a.json", "N1", "live", "2026-09-01T00:00:00+00:00", recall=(7, 8))
    _write_node(tmp_path / "b.json", "N1", "gold", "2026-09-28T00:00:00+00:00")
    _write_node(tmp_path / "c.json", "N2", "gold", "2026-09-30T00:00:00+00:00")
    _write_node(tmp_path / "d.json", "N2", "live", "2026-09-02T00:00:00+00:00", recall=(1, 9))
    _write_node(tmp_path / "e.json", "N2", "live", "2026-09-29T00:00:00+00:00", recall=(5, 9))
    _write_node(tmp_path / "f.json", "N3", "gold", "2026-09-29T00:00:00+00:00")
    replays = load_replays(tmp_path)
    assert set(replays) == {"N1", "N2", "N3"}
    assert replays["N1"].meta["source"] == "live"
    assert (replays["N2"].meta["source"], replays["N2"].meta["captured_at"]) == (
        "live", "2026-09-29T00:00:00+00:00"
    )
    assert replays["N3"].is_gold
    assert load_replay(tmp_path).meta["node"] == "N1"  # 기존 API = load_replays()["N1"]


# (b) 미캡처 노드 — 결과에 없고, 표에 행이 없고, 합계는 "측정 노드 없음".
def test_uncaptured_nodes_are_absent_and_not_counted(tmp_path: Path) -> None:
    _write_node(tmp_path / "n1.json", "N1", "live", "2026-09-29T00:00:00+00:00", recall=(7, 8))
    _write_node(tmp_path / "n3.json", "N3", "gold", "2026-09-29T00:00:00+00:00")
    replays = load_replays(tmp_path)
    assert "N2" not in replays and "N4" not in replays
    table = service.evaluation_table(replays)
    assert [r["노드"].split()[0] for r in table[:-1]] == ["N1", "N3"]
    assert table[1]["recall(m/n)"] == "골드 재생 — 해당 없음"  # 골드 재생은 측정이 아니다
    assert "측정 노드 없음" in table[-1]["노드"] and table[-1]["recall(m/n)"] == "—"
    (tmp_path / "n1.json").unlink()
    with pytest.raises(FileNotFoundError):
        load_replay(tmp_path)  # N1 이 없으면 기존 API 는 명시적으로 실패한다


# (c) 홀드아웃 합계 recall = Σmatched / Σgold — 세 노드면 분모 26, 일부면 실제 분모와 노드 표기.
def test_holdout_total_recall_denominator(tmp_path: Path) -> None:
    _write_node(tmp_path / "n1.json", "N1", "live", "2026-09-29T00:00:00+00:00", recall=(7, 8))
    _write_node(tmp_path / "n2.json", "N2", "live", "2026-09-29T00:00:00+00:00", recall=(5, 9))
    _write_node(tmp_path / "n3.json", "N3", "live", "2026-09-29T00:00:00+00:00", recall=(3, 7))
    _write_node(tmp_path / "n4.json", "N4", "live", "2026-09-29T00:00:00+00:00", recall=(4, 10))
    total = service.evaluation_table(load_replays(tmp_path))[-1]
    assert total["노드"] == "홀드아웃 합계"
    assert total["recall(m/n)"] == f"{12 / 26:.3f} (12/26)"  # N1(tune) 7/8 은 섞이지 않는다
    (tmp_path / "n4.json").unlink()
    total = service.evaluation_table(load_replays(tmp_path))[-1]
    assert total["노드"] == "홀드아웃 합계 (N2·N3 만 — 골드 16/26건)"
    assert total["recall(m/n)"] == "0.500 (8/16)"
    markdown = service.evaluation_markdown(load_replays(tmp_path))
    assert len(markdown.splitlines()) == 2 + 3 + 1  # 머리·구분 + N1·N2·N3 + 합계
    assert "| 0.500 (8/16) |" in markdown


# (d) 9/29 N1 파일의 `recall_n1` 키와 I-1 의 `recall` 키를 둘 다 읽는다.
@pytest.mark.parametrize("key", ["recall_n1", "recall"])
def test_recall_key_compat(tmp_path: Path, key: str) -> None:
    _write_node(tmp_path / "n1.json", "N1", "live", "2026-09-29T00:00:00+00:00", recall=(7, 8), recall_key=key)
    result = load_replay(tmp_path)
    assert result.meta["recall"] == {"recall": 0.875, "matched": 7, "total": 8}
    assert "recall_n1" not in result.meta
    assert "N1(tune) recall 0.875 (7/8)" in service.summary_line(result)


def test_committed_n1_replay_reads_recall_n1_key() -> None:
    assert load_replay().meta["recall"]["total"] == 8


# (e) AppTest — 공정 3개 × 노드 전환(J-04). 예외 0, 미캡처는 비활성, 평가 표는 접힌 expander 안.
def test_app_switches_between_processes_and_nodes(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    monkeypatch.delenv("HAZOP_ALLOW_LIVE", raising=False)
    replays = load_replays()
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    assert len(at.table) == 4  # 안내 탭 3(가이드워드·S·F·위험도 구간) + 평가 요약 1
    assert next(e for e in at.expander if e.label.startswith("정확도")).proto.expanded is False
    at.radio(key="mode").set_value("실측 사례 재생").run()
    assert at.selectbox(key="process_name").options == [p["name"] for p in service.CATALOG]
    for process in service.CATALOG:
        at.selectbox(key="process_name").select(process["name"]).run()
        assert not at.exception
        for node in process["nodes"]:
            nid = node["id"]
            button = at.button(key=f"preset_{nid}")
            assert button.disabled is (nid not in replays)
            assert ("미캡처" in button.label) is (nid not in replays)
            if nid not in replays:
                continue
            button.click().run()
            assert not at.exception
            assert len(at.dataframe) == 1
            assert len(at.dataframe[0].value) == len(replays[nid].records)
            assert any(nid in m.value for m in at.markdown)


# ── self-verification T-03 ──────────────────────────────────────────────────
# 결함 1건 삽입 → 화면 표와 xlsx 신뢰도 시트가 모두 review, 요약 줄에 건수, 파일 원본은 불변.
def test_verifier_flag_reaches_table_and_xlsx() -> None:
    replay = load_replays()["N2"]
    records = list(replay.records)
    records[2] = records[2].model_copy(update={"causes": [*records[2].causes, "API 520 기준 릴리프 용량 부족"]})
    result = service.Result(meta=dict(replay.meta), records=records)

    table = service.worksheet_table(result)
    reviewed = [i for i, row in enumerate(table) if "🔴" in str(row[service.CONFIDENCE_COLUMN])]
    assert reviewed == [2]
    assert table[2][service.FLAG_COLUMN] == "unverified_standard: API 520"
    assert all(row[service.FLAG_COLUMN] == "" for i, row in enumerate(table) if i != 2)
    assert "review 1건 (규격 1·수치 0)" in service.summary_line(result)

    ws = openpyxl.load_workbook(io.BytesIO(service.export_files(result)["xlsx"][1]))["신뢰도"]
    labels = [r[1] for r in ws.iter_rows(min_row=2, values_only=True)]
    review_label, _ = service.confidence_label("review")
    assert [i for i, v in enumerate(labels) if v == review_label] == [2]
    assert records[2].confidence == "inferred"  # 원본 레코드는 격하되지 않는다
    assert result.verified is not None  # 표·내보내기가 같은 캐시를 썼다
