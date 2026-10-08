"""지시문 Y-2 평가기준 다변화 — 기준 파일·위험도 산정·생성·내보내기 (오프라인)."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from core.agent import HazopGenerator, NodeMeta
from core.agent.generate import DEVIATION_BATCH_SCHEMA, batch_schema
from core.criteria import GOLD_CRITERIA, OFFICIAL_CRITERIA, all_criteria, load_criteria
from core.export import build_report, export_all, normalize_rows
from core.llm import ConverseResponse, Message, MockBedrockClient

_ROOT = Path(__file__).parent.parent
#: C-C-37-2026 <표 2> 위험도 대조표 예시(인쇄 9쪽) — 행 강도 4→1, 열 발생빈도 3(상)·2(중)·1(하). 원문 배치 그대로 옮겨 적었다.
_CC37_TABLE2 = {4: (5, 5, 3), 3: (4, 4, 2), 2: (3, 2, 1), 1: (2, 1, 1)}
_META = NodeMeta(node="X1", substance="수소", phase="gas", equipment=["압축기", "디스펜서"], safeguards=[])


def test_cc37_matches_table2_and_keeps_original_levels() -> None:
    c = load_criteria(OFFICIAL_CRITERIA)
    assert (c.s_max, c.f_max, c.method) == (4, 3, "lookup")  # 5단계로 바꾸지 않는다
    for s, row in _CC37_TABLE2.items():
        assert [c.risk(s, f) for f in (3, 2, 1)] == list(row)
    assert [c.band(x) for x in (1, 2, 3, 4, 5)] == ["위험작업 수용"] * 2 + ["조건부 위험작업 수용"] * 2 + ["위험작업 불허"]


@pytest.mark.parametrize("criteria_id", [c.id for c in all_criteria()])
def test_excel_formula_equals_risk_for_every_cell(criteria_id: str) -> None:
    """xlsx 위험도 수식을 손으로 풀어 코드 값과 대조한다 — 검토에서 S·F 를 고쳐도 같은 값이 나와야 한다."""
    c = load_criteria(criteria_id)
    formula = c.excel_formula(2)
    for s in range(1, c.s_max + 1):
        for f in range(1, c.f_max + 1):
            if formula == "=H2*I2":
                value = s * f
            else:
                matrix = re.fullmatch(r"=INDEX\(\{(.+)\},H2,I2\)", formula).group(1)  # type: ignore[union-attr]
                value = int(matrix.split(";")[s - 1].split(",")[f - 1])
            assert value == c.risk(s, f), (criteria_id, s, f)


@pytest.mark.parametrize("criteria_id", [c.id for c in all_criteria()])
def test_every_criteria_file_carries_source_and_license(criteria_id: str) -> None:
    data = load_criteria(criteria_id).data
    for key in ("name", "issuer", "doc", "locator", "license", "severity", "frequency", "risk_bands"):
        assert data.get(key), (criteria_id, key)
    if criteria_id != GOLD_CRITERIA:
        assert str(data["url"]).startswith("https://") and data["retrieved_at"]
    bands = load_criteria(criteria_id).bands()
    covered = sorted(x for low, high, _, _ in bands for x in range(low, high + 1))
    c = load_criteria(criteria_id)
    risks = {c.risk(s, f) for s in range(1, c.s_max + 1) for f in range(1, c.f_max + 1)}
    assert risks <= set(covered), "위험도 값 중 구간에 안 들어가는 것이 있다"


def test_gold_prompt_text_is_byte_identical_to_rating_scale() -> None:
    """골드셋 기준으로 부르면 판정 프롬프트가 Y-2 이전과 같다 — 기존 실측(recall·비용)이 그대로 유효."""
    raw = (_ROOT / "data" / "gold" / "rating_scale.json").read_text(encoding="utf-8").strip()
    assert load_criteria(GOLD_CRITERIA).prompt_text() == raw
    assert load_criteria(None).id == GOLD_CRITERIA
    assert batch_schema(load_criteria(GOLD_CRITERIA)) == DEVIATION_BATCH_SCHEMA


def _batch(guideword: str, s: int, f: int) -> str:
    cells = [{"parameter": p, "applicable": True, "deviation": f"{p} 이탈", "causes": ["c"], "consequences": ["r"],
              "safeguards_before": [], "S": s, "F": f, "recommendations": ["k"], "evidence": [], "confidence": "inferred"}
             for p in ("유량", "압력", "온도", "준위", "조성", "상")]
    return json.dumps({"guideword": guideword, "cells": cells}, ensure_ascii=False)


def _generator(s: int, f: int, criteria_id: str) -> tuple[HazopGenerator, MockBedrockClient]:
    def make(system: str, messages: list[Message], **_: Any) -> ConverseResponse:
        if "파라미터 축" in system:
            names = ["유량", "압력", "온도", "준위", "조성", "상"]
            return ConverseResponse(content=json.dumps({"parameters": [{"name": n, "rationale": "r"} for n in names]}))
        gw = next(g for g in ("No", "More", "Less", "Reverse", "Other than", "Part of", "As well as")
                  if f"\n{g} —" in str(messages[0].content))
        return ConverseResponse(content=_batch(gw, s, f))

    client = MockBedrockClient(response_factory=make)
    return HazopGenerator(client, criteria_id=criteria_id), client


def test_official_generation_uses_cc37_prompt_schema_and_lookup_risk() -> None:
    gen, client = _generator(4, 1, OFFICIAL_CRITERIA)
    records = gen.generate(_META)
    assert records and {r.criteria_id for r in records} == {OFFICIAL_CRITERIA}
    assert {r.risk_score for r in records} == {3}  # 대조표 S4·F1 = 3 (곱이면 4)
    judge = client.calls[1]
    assert "C-C-37-2026" in judge["system"] and "S 는 1~4 정수, F 는 1~3 정수" in judge["system"]
    cell = judge["response_schema"]["properties"]["cells"]["items"]["properties"]
    assert (cell["S"]["maximum"], cell["F"]["maximum"]) == (4, 3)


def test_out_of_range_grade_is_rejected_not_clamped() -> None:
    """모델이 NH3 습관대로 F=5 를 내면 C-C-37 스키마가 막는다 → 그 가이드워드는 review(조용히 자르지 않는다)."""
    gen, _ = _generator(3, 5, OFFICIAL_CRITERIA)
    assert gen.generate(_META) == []
    assert len(gen.review_guidewords) == 7


def test_report_and_xlsx_follow_record_criteria() -> None:
    gen, _ = _generator(3, 2, OFFICIAL_CRITERIA)
    records = gen.generate(_META)
    report = build_report(normalize_rows(records))
    assert report.criteria_id == OFFICIAL_CRITERIA
    assert list(report.risk_distribution) == ["위험작업 불허", "조건부 위험작업 수용", "위험작업 수용"]
    assert report.risk_distribution["조건부 위험작업 수용"] == len(records)  # S3·F2 = 4
    assert set(report.sf_matrix) == {"1", "2", "3", "4"} and set(report.sf_matrix["1"]) == {"1", "2", "3"}


def test_xlsx_headers_rating_sheet_and_screening_size(tmp_path: Path) -> None:
    gen, _ = _generator(2, 3, OFFICIAL_CRITERIA)
    paths = export_all(gen.generate(_META), tmp_path)
    wb = load_workbook(io.BytesIO(paths["xlsx"].read_bytes()))
    head = [c.value for c in wb["HAZOP워크시트"][1]]
    assert head[7:9] == ["S(1-4)", "F(1-3)"]
    rating = [tuple(r)[:3] for r in wb["평가기준"].iter_rows(values_only=True)]
    assert ("S 강도", 4, "치명적: 사망, 부상 2명 이상, 재산손실 10억원 이상, 설비 운전정지 기간 10일 이상") in rating
    assert any(r[0] == "출처" and "C-C-37-2026" in str(r[2]) for r in rating)
    assert any(r[0] == "이용 조건" for r in rating)
    screening = list(wb["스크리닝"].iter_rows(values_only=True))
    assert screening[0][0].startswith("4×3") and screening[1] == ("S＼F", "F1", "F2", "F3")
    assert "대조표" in paths["lopa"].read_text(encoding="utf-8")


def test_mixed_criteria_in_one_export_is_refused(tmp_path: Path) -> None:
    gold = json.loads((_ROOT / "data" / "gold" / "hazop_nh3.json").read_text(encoding="utf-8"))[:2]
    gen, _ = _generator(2, 2, OFFICIAL_CRITERIA)
    with pytest.raises(ValueError, match="섞였다"):
        export_all([*gold, *gen.generate(_META)], tmp_path)
