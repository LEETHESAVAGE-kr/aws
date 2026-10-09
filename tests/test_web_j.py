"""오프라인 시험 — FR-10 J-01·J-03·J-04 (지시문 J): 공정 카탈로그·직접 입력 빠른 실호출·화면 재배치. 네트워크 0회."""

from __future__ import annotations

import json
from datetime import date
from typing import TYPE_CHECKING, Any

import pytest

from apps.web import service
from apps.web.catalog import CATALOG_PATH, load_catalog
from apps.web.replay import REPLAY_DIR, load_replays
from core.agent import NodeMeta
from core.export.xlsx import REVIEW_HEADERS, REVIEW_SHEET, SHEET_ORDER

if TYPE_CHECKING:
    from collections.abc import Callable
    from pathlib import Path

_APP = REPLAY_DIR.parents[1] / "apps" / "web" / "app.py"
_MOCK_ENV = {"HAZOP_USE_MOCK": "true", "HAZOP_ALLOW_LIVE": "true"}
_CASES = "실측 사례 재생"
_GENERATE = "HAZOP 초안 생성"
_SENTENCE = "수소충전소 압축기에서 디스펜서로 고압 수소를 보낸다. 안전장치는 긴급차단밸브."


@pytest.fixture(autouse=True)
def _reset_daily_counter() -> None:
    service._daily_runs.clear()


# ── J-01 공정 카탈로그 ─────────────────────────────────────────────────────────
def test_catalog_loads_validated_unique_nodes() -> None:
    catalog = load_catalog()
    assert [p["id"] for p in catalog] == ["nh3_sts", "lpg_loading", "cl2_unloading", "lpg_iocl"]
    assert [p["gold"] for p in catalog] == [True, False, False, False]  # 공개 HAZOP 대조는 골드셋이 아니다
    ids = [n["id"] for p in catalog for n in p["nodes"]]
    assert ids == ["N1", "N2", "N3", "N4", "P1", "P2", "I1", "I2", "I3", "I4"]
    assert all(isinstance(n["node_meta"], NodeMeta) for p in catalog for n in p["nodes"])
    # NH3 는 I-1 의 하드코딩 값 그대로(J-01 "추측 금지") — PRESETS 도 예전 dict 와 같다.
    assert {k: service.PRESETS[k] for k in ("N1", "N2", "N3", "N4")} == {
        "N1": "벙커링선 매니폴드", "N2": "이송 호스", "N3": "수급선 매니폴드", "N4": "이송 운전 절차",
    }
    for node in catalog[0]["nodes"]:
        m = node["node_meta"]
        assert (m.substance, m.phase, m.P_kPag, m.T_degC, m.safeguards) == ("NH3", "unknown", None, None, [])


def _dup(c: dict[str, Any]) -> None:
    c["processes"][1]["nodes"][0]["id"] = "N1"


def _bad_phase(c: dict[str, Any]) -> None:
    c["processes"][1]["nodes"][0]["node_meta"]["phase"] = "액상"


def _id_mismatch(c: dict[str, Any]) -> None:
    c["processes"][2]["nodes"][0]["node_meta"]["node"] = "P9"


@pytest.mark.parametrize(("mutate", "message"), [(_dup, "중복"), (_bad_phase, "스키마"), (_id_mismatch, "불일치")])
def test_catalog_rejects_bad_entries(
    tmp_path: Path, mutate: Callable[[dict[str, Any]], None], message: str
) -> None:
    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    mutate(catalog)
    path = tmp_path / "presets.json"
    path.write_text(json.dumps(catalog, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        load_catalog(path)


def test_capture_non_gold_node_saves_null_recall(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from tools import capture_replay

    assert capture_replay.default_split("P1") == "none"
    assert capture_replay.default_split("N2") == "holdout"
    client = service.MockBedrockClient(response_factory=service._mock_factory(load_replays()["N1"]))
    monkeypatch.setattr(capture_replay, "get_bedrock_client", lambda: client)
    out = tmp_path / "p1.json"
    # --split 을 줘도 골드 없는 노드는 none 으로 고정된다.
    assert capture_replay.main(["--node", "P1", "--out", str(out), "--source", "live", "--split", "holdout"]) == 0
    saved = json.loads(out.read_text(encoding="utf-8"))
    assert (saved["node"], saved["split"], saved["recall"]) == ("P1", "none", None)
    assert saved["records"] and {r["node_meta"]["substance"] for r in saved["records"]} == {"프로판"}
    assert saved["parameters"] and len(saved["raw_calls"]) == len(client.calls)
    # 예시 공정 재생은 recall 표에 행을 만들지 않는다.
    replays = load_replays(tmp_path)
    assert "P1" in replays and service.evaluation_table(replays)[:-1] == []
    with pytest.raises(ValueError, match="골드셋이 없는"):
        capture_replay.main(["--node", "P1", "--out", str(out), "--source", "gold"])


# ── J-03 직접 입력 빠른 실호출 ─────────────────────────────────────────────────
def test_run_quick_makes_two_calls_for_one_guideword() -> None:
    node = {**service.NODES["P1"]["node_meta"].model_dump(), "node": "X1"}
    result = service.run_quick(json.dumps(node, ensure_ascii=False), "More", load_replays()["N1"], _MOCK_ENV)
    m = result.meta
    assert len(m["raw_calls"]) == 2  # 열거 1 + 가이드워드 1
    assert result.records and {r.guideword for r in result.records} == {"More"}
    assert (m["source"], m["mock"], m["guidewords"], m["recall"]) == ("quick", True, ["More"], None)
    assert m["expected_cells"] == len(m["parameters"]) == m["judged_cells"]
    assert {r.node for r in result.records} == {"X1"}
    assert "X1 정성 검토용 — 대조 기준 없음" in service.summary_line(result)
    assert set(service.export_files(result)) == {"xlsx", "lopa", "report"}
    assert service.process_view(result)["api_calls"] == 2


@pytest.mark.parametrize(
    ("node", "guideword"),
    [
        ({"substance": "프로판", "phase": "액상", "P_kPag": 1, "T_degC": 1, "equipment": [], "safeguards": []}, "More"),
        ({"substance": "프로판"}, "More"),
        ({"substance": "프로판", "phase": "gas", "P_kPag": 1, "T_degC": 1, "equipment": [], "safeguards": []}, "Much"),
    ],
)
def test_run_quick_rejects_bad_input_before_any_call(
    node: dict[str, Any], guideword: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _boom(*_: object, **__: object) -> None:
        raise AssertionError("API 를 부르면 안 된다")

    monkeypatch.setattr(service, "HazopGenerator", _boom)
    with pytest.raises(ValueError):
        service.run_quick(json.dumps(node, ensure_ascii=False), guideword, load_replays()["N1"], _MOCK_ENV)


# ── 자연어 입력 (사용자 결정 9/29, PRD FR-10 밖) ──────────────────────────────
def _parse_client(payload: dict[str, Any]) -> service.MockBedrockClient:
    return service.MockBedrockClient(
        response_factory=lambda **_: service.ConverseResponse(content=json.dumps(payload, ensure_ascii=False))
    )


def test_run_quick_natural_language_adds_one_parse_call() -> None:
    result = service.run_quick(_SENTENCE, "More", load_replays()["N1"], _MOCK_ENV)
    m = result.meta
    assert [c.get("stage") for c in m["raw_calls"]] == ["parse", None, None]  # 해석 1 + 열거 1 + 판정 1
    assert (m["parsed_by"], m["node_text"], m["parse_model"]) == ("llm", _SENTENCE, "mock")
    assert m["node_meta"]["substance"] == "수소" and {r.guideword for r in result.records} == {"More"}


def test_parse_json_input_makes_no_call() -> None:
    client = _parse_client({})
    node = service.NODES["P2"]["node_meta"].model_dump()
    meta, parsed = service.parse_node_text(json.dumps(node, ensure_ascii=False), _MOCK_ENV, client)
    assert meta.substance == "염소" and parsed["parsed_by"] == "json" and client.calls == []


def test_parse_sends_sentence_with_schema_and_validates_reply() -> None:
    reply = {"node": "X1", "substance": "수소", "phase": "gas", "pressure": {"value": 90, "unit": "MPa"},
             "T_degC": None, "equipment": ["압축기"], "safeguards": [],
             "design_pressure": {"value": 100, "unit": "MPa"}, "capacity": "200 kg"}
    client = _parse_client(reply)
    meta, parsed = service.parse_node_text(_SENTENCE, _MOCK_ENV, client)
    assert meta.P_kPag == 90000 and meta.T_degC is None and parsed["cost_usd"] == 0.0  # 환산은 코드가 한다
    assert meta.design_P_kPag == 100000 and meta.capacity == "200 kg"  # 10/9 설계압력·용량 칸
    (call,) = client.calls
    assert _SENTENCE in str(call["messages"][0].content) and "{text}" not in str(call["messages"][0].content)
    schema = call["response_schema"]
    assert schema["properties"]["phase"]["enum"] == ["liquid", "gas", "liquid/gas", "unknown"]
    assert "P_kPag" not in schema["properties"] and "pressure" in schema["required"]
    assert {"design_pressure", "capacity"} <= set(schema["required"]) and "design_P_kPag" not in schema["properties"]
    assert "node" in schema["required"] and "$comment" not in json.dumps(schema)
    # 모델이 enum 밖 값을 내면 core/llm 이 응답 스키마로 1회 재시도 후 content=None → 생성 전에 막힌다.
    # (그 뒤의 validate_node_meta 는 _parse_schema 와 산출물 스키마가 어긋날 때를 위한 2차 방어라 여기선 닿지 않는다.)
    bad = _parse_client({**reply, "phase": "기체"})
    with pytest.raises(ValueError, match="2회 위반"):
        service.parse_node_text(_SENTENCE, _MOCK_ENV, bad)
    assert len(bad.calls) == 2


@pytest.mark.parametrize(
    ("pressure", "kpag"),
    [(None, None), ({"value": 7, "unit": "bar"}, 700), ({"value": 350, "unit": "kPa"}, 350),
     ({"value": 10, "unit": "kgf/cm2"}, 980.665)],
)
def test_parse_converts_pressure_in_code(pressure: dict[str, Any] | None, kpag: float | None) -> None:
    reply = {"node": "X1", "substance": "프로판", "phase": "liquid", "pressure": pressure, "T_degC": 25,
             "equipment": ["저장탱크"], "safeguards": [], "design_pressure": None, "capacity": None}
    meta, _ = service.parse_node_text(_SENTENCE, _MOCK_ENV, _parse_client(reply))
    assert meta.P_kPag == kpag


@pytest.mark.parametrize("text", ["", "   ", "가" * (service.NODE_TEXT_LIMIT + 1)])
def test_parse_rejects_empty_or_long_text_before_call(text: str) -> None:
    client = _parse_client({})
    with pytest.raises(ValueError):
        service.parse_node_text(text, _MOCK_ENV, client)
    assert client.calls == []


def test_parser_uses_low_cost_verifier_profile(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")  # 생성자만 — 호출하지 않는다
    config = service.load_model_config()
    assert service._parser_client().model_short == config.verifier.model_id != config.generation.model_id


@pytest.mark.parametrize(
    ("value", "scope"), [(None, "quick"), ("quick", "quick"), ("FULL", "full"), ("everything", "quick")]
)
def test_live_scope(value: str | None, scope: str) -> None:
    assert service.live_scope({} if value is None else {"HAZOP_LIVE_SCOPE": value}) == scope


# ── J-04 화면 — 직접 입력 mock 경로·상한 공용·HAZOP_LIVE_SCOPE 분기 ─────────────
def _guidewords_in(frame: Any) -> set[str]:
    return {v.split(" (")[0] for v in frame.value["가이드워드"]}  # "More (압력)" 형식


@pytest.mark.parametrize("scope", ["quick", "full"])  # 옛 HAZOP_LIVE_SCOPE 값은 화면을 바꾸지 않는다(X-1)
def test_app_default_run_covers_all_guidewords(monkeypatch: pytest.MonkeyPatch, scope: str) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in {**_MOCK_ENV, "HAZOP_LIVE_SCOPE": scope}.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=60).run()
    assert not at.exception
    assert len(at.dataframe) == 0
    assert not any("노드 전체" in b.label for b in at.button)
    assert not any("어떤 이탈" in m.value for m in at.markdown)  # X-G5 — 선택 질문 블록이 없다
    # 기본 = 문장 모드. 입력 칸은 비어 있고 예시는 칩이 맡는다 — 비어 있으면 버튼이 막힌다.
    assert at.radio(key="mode").value == "문장으로 새 공정 분석"
    assert at.text_area[0].value == "" and "예시" in at.text_area[0].placeholder
    assert _cta(at).label == service.LIVE_BUTTON and _cta(at).disabled is True
    at.text_area[0].input(_SENTENCE).run()
    assert _cta(at).disabled is False
    _cta(at).click().run()
    assert not at.exception
    assert len(at.dataframe) == 1
    from core.agent.generate import STANDARD_GUIDEWORDS

    assert _guidewords_in(at.dataframe[0]) == set(STANDARD_GUIDEWORDS)  # X-G1 표준 7종 전부
    assert any("API 호출 9회" in m.value for m in at.markdown)  # 해석 1 + 열거 1 + 판정 7
    assert any(_SENTENCE in c.value for c in at.code)  # 입력 문장이 해석 결과 옆에 보인다
    assert not any("1종" in m.value for m in at.metric)
    # 상한은 공용 — 1회 뒤엔 보조 실행 버튼도 막히고, 사례 모드로 안내한다.
    assert _cta(at).disabled and at.button(key="run_quick").disabled
    assert any("세션" in c.value for c in at.caption)
    next(b for b in at.button if b.label == "실측 사례 보기").click().run()
    assert not at.exception and at.radio(key="mode").value == _CASES


def test_app_quick_run_is_one_guideword_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    """X-1b 보조 실행: 고른 가이드워드 1종만, 요약 줄·출처 줄·판정 셀 타일에 '1종' 이 보인다."""
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert at.expander[0].label == service.QUICK_EXPANDER and at.selectbox(key="guideword").label == (
        service.QUICK_QUESTION
    )
    at.text_area[0].input(_SENTENCE).run()
    at.selectbox(key="guideword").select("Less").run()
    at.button(key="run_quick").click().run()
    assert not at.exception
    assert _guidewords_in(at.dataframe[0]) == {"Less"}
    assert any("API 호출 3회" in m.value for m in at.markdown)  # 해석 1 + 열거 1 + 판정 1
    scope = service.QUICK_SCOPE.format(guideword="Less")
    assert any(scope in c.value for c in at.caption)  # 요약 줄·출처 줄
    assert any(m.value.endswith("(1종)") for m in at.metric)


def test_quick_feature_copy_has_no_booth_words() -> None:
    """X-1b 금지어 — 보조 실행은 실무 기능이다. 이 기능의 문구에 '부스'·'관람객' 이 없다(사용자 결정 10/8)."""
    quick = service.run_quick(_SENTENCE, "More", load_replays()["N1"], _MOCK_ENV)
    texts = [service.QUICK_EXPANDER, service.QUICK_QUESTION, service.QUICK_BUTTON, service.QUICK_SCOPE,
             service.summary_line(quick), service.provenance_line(quick)]
    for text in texts:
        assert "부스" not in text and "관람객" not in text, text


def test_app_direct_input_disabled_without_allow(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    monkeypatch.delenv("HAZOP_ALLOW_LIVE", raising=False)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    assert [b.disabled for b in at.button if b.label.startswith(_GENERATE)] == [True]
    assert all(not b.disabled for b in at.button if (b.key or "").startswith("chip_"))
    assert any("HAZOP_ALLOW_LIVE" in c.value for c in at.caption)


# ── 시크릿 공백 정리 · 401 키 진단 (9/29 배포 401) ──────────────────────────────
_FAKE_KEY = "sk-ant-api03-" + "A1b2_C3d4-" * 9 + "xyz"  # 형식만 흉내 낸 가짜(106자)


def test_sync_secrets_strips_whitespace_new_and_existing() -> None:
    environ: dict[str, str] = {"HAZOP_LIVE_SCOPE": " full\n"}
    service.sync_secrets({"ANTHROPIC_API_KEY": f"  {_FAKE_KEY}\n", "HAZOP_ALLOW_LIVE": "true "}, environ)
    assert environ == {"ANTHROPIC_API_KEY": _FAKE_KEY, "HAZOP_ALLOW_LIVE": "true", "HAZOP_LIVE_SCOPE": "full"}
    service.sync_secrets({"ANTHROPIC_API_KEY": "other"}, environ)
    assert environ["ANTHROPIC_API_KEY"] == _FAKE_KEY  # 이미 있는 값은 덮어쓰지 않는다


@pytest.mark.parametrize(
    ("key", "expected"),
    [
        (_FAKE_KEY, ["시작 ✅", f"길이 {len(_FAKE_KEY)}자", "외 문자 없음 ✅"]),
        ("sk-ant-...", ["예시 값", "외 문자 3개 ❌"]),
        ("apikey_1234", ["시작하지 않음 ❌", "길이 11자"]),
        (f"\u201c{_FAKE_KEY}\u201d", ["시작하지 않음 ❌", "외 문자 2개 ❌"]),  # 스마트 따옴표
        (f"{_FAKE_KEY} \n", ["시작 ✅", "공백·줄바꿈 있음"]),
        ("", ["비어 있습니다"]),
    ],
)
def test_key_hint_never_reveals_key(key: str, expected: list[str]) -> None:
    hint = service.key_hint({"ANTHROPIC_API_KEY": key})
    assert all(part in hint for part in expected), hint
    assert "api03" not in hint and "A1b2" not in hint  # 접두사 'sk-ant-' 외 키 문자는 한 글자도 없다


class AuthenticationError(Exception):
    """anthropic.AuthenticationError 대역 — 이름으로 판정되는지 본다."""


class _Http401Error(Exception):
    status_code = 401


def test_is_auth_error() -> None:
    assert service.is_auth_error(AuthenticationError("401"))
    assert service.is_auth_error(_Http401Error())
    assert not service.is_auth_error(ValueError("x"))


def test_app_shows_key_hint_on_auth_error(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in {**_MOCK_ENV, "ANTHROPIC_API_KEY": _FAKE_KEY}.items():
        monkeypatch.setenv(key, value)

    def _reject(*_: object, **__: object) -> None:
        raise AuthenticationError("Error code: 401 - API key is invalid.")

    monkeypatch.setattr(service, "run_live", _reject)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    at.text_area[0].input(_SENTENCE).run()
    next(b for b in at.button if b.label.startswith(_GENERATE)).click().run()
    assert not at.exception
    assert any("AuthenticationError" in e.value for e in at.error)
    hints = [c.value for c in at.caption if c.value.startswith("키 진단")]
    assert hints and "시작 ✅" in hints[0] and "api03" not in hints[0]


# ── M-03 평가기준 불일치 배지 · F 분포 (실무자평가 P-2·P-3) ──────────────────────
def test_criteria_notice_only_for_non_gold_processes() -> None:
    replays = load_replays()
    assert service.criteria_notice(replays["N1"]) is None
    assert service.criteria_notice(replays["N4"]) is None
    assert service.criteria_notice(replays["P1"]) == service.CRITERIA_NOTICE  # 10/9 이전 캡처 — NH3 기준으로 생성됨
    quick = service.run_quick(_SENTENCE, "More", replays["N1"], _MOCK_ENV)
    # Y-2: 직접 입력은 공식 기준(C-C-37)으로 매기고 그 이름·위치를 적는다 — '참고용' 배지가 아니다
    notice = service.criteria_notice(quick)
    assert notice != service.CRITERIA_NOTICE and "C-C-37-2026" in notice and "6.5" in notice
    assert {r.criteria_id for r in quick.records} == {"kosha_cc37_2026"}
    assert all(1 <= r.S <= 4 and 1 <= r.F <= 3 for r in quick.records)


@pytest.mark.parametrize(
    ("node", "expected"),
    [("N1", "빈도=3 비율 93%"), ("N4", "빈도=3 비율 95%"), ("P1", "빈도=3 비율 85%"), ("P2", "빈도=3 비율 64%")],
)
def test_f_distribution_from_replays(node: str, expected: str) -> None:
    """재생 파일의 F 분포가 그대로 나온다 — 불리한 숫자도 그대로. N4 는 실무자평가 P-2 값(92/97),
    N1·P1·P2 는 O-2 재캡처 값(68/73·61/72·45/70). 1차 평가 값(52/61·60/69·54/64)은 진행로그 대조표에."""
    result = load_replays()[node]
    assert service.f_distribution(result) == expected
    assert expected in service.summary_line(result)


def test_f_distribution_tie_and_empty() -> None:
    base = load_replays()["N1"]
    two = [base.records[0].model_copy(update={"F": 4}), base.records[0].model_copy(update={"F": 2})]
    assert service.f_distribution(service.Result(meta={}, records=two)) == "빈도=2 비율 50%"
    assert service.f_distribution(service.Result(meta={}, records=[])) == "빈도 분포 해당 없음"


def test_app_criteria_badge_visibility(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)

    def badge_shown(at: AppTest) -> bool:
        return any(service.CRITERIA_NOTICE in m.value for m in at.markdown)

    def official_shown(at: AppTest) -> bool:
        return any("C-C-37-2026" in m.value and "평가기준" in m.value for m in at.markdown)

    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    at.radio(key="mode").set_value(_CASES).run()
    assert not at.exception and len(at.dataframe) == 1 and not badge_shown(at)  # 사례 첫 화면 = N1(골드 공정)
    at.selectbox(key="process_name").select("LPG 저장탱크 출하").run()
    assert not at.exception and badge_shown(at)
    at.radio(key="mode").set_value("문장으로 새 공정 분석").run()
    at.text_area[0].input(_SENTENCE).run()
    next(b for b in at.button if b.label.startswith(_GENERATE)).click().run()
    assert not at.exception and len(at.dataframe) == 1 and not badge_shown(at) and official_shown(at)  # Y-2


def test_provenance_marks_replays_captured_before_m01() -> None:
    """홀드아웃 N2~N4 재생은 M-01 이전 프롬프트 결과라 출처 줄에 그렇게 적힌다. O-2 재캡처(P1·P2)는 이후라 안 붙고
    병렬 수가 붙는다."""
    replays = load_replays()
    for node in ("N2", "N3", "N4"):
        assert "M-01 이전 프롬프트" in service.provenance_line(replays[node]), node
    for node in ("P1", "P2"):
        line = service.provenance_line(replays[node])
        assert "M-01 이전" not in line and " · 병렬 4" in line, (node, line)
    later = service.Result(meta={**replays["P1"].meta, "captured_at": "2026-09-29T05:00:00+00:00"})
    assert "M-01 이전" not in service.provenance_line(later)


# ── LOPA 초안 Word 다운로드 (사용자 결정 9/29) ─────────────────────────────────
def test_lopa_download_is_docx_with_markdown_content_preserved(monkeypatch: pytest.MonkeyPatch) -> None:
    """파일은 진짜 .docx 이고, 원본 Markdown 의 모든 줄 글자가 문단·표 칸에 그대로 들어간다."""
    import io
    import re

    from docx import Document

    from apps.web import docx_export
    from core.export import lopa

    seen: list[str] = []
    convert = docx_export.lopa_markdown_to_docx
    monkeypatch.setattr(service, "lopa_markdown_to_docx", lambda md: (seen.append(md), convert(md))[1])
    name, data = service.export_files(load_replays()["P1"])["lopa"]
    assert name == "lopa_draft.docx" and data[:2] == b"PK" and len(seen) == 1
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert lopa.DISCLAIMER in paragraphs
    assert sum(bool(re.match(r"시나리오 \d+:", p)) for p in paragraphs) == 5
    assert all(row in cells for row in lopa.QUANT_ROWS)
    checked = 0
    for line in seen[0].splitlines():
        if not line.strip() or re.match(r"^\|[\s:|-]+\|$", line.strip()):
            continue
        if line.startswith("|"):
            assert all(c.replace("`", "") in cells for c in docx_export._cells(line))
        else:
            assert re.sub(r"^(#{1,3} |> |- )", "", line).replace("**", "").replace("`", "") in paragraphs, line
        checked += 1
    assert checked > 50
    assert not any("`" in p for p in paragraphs + cells)


# ── 화면 문구 개편 (사용자 요청 9/29) — 정확도 줄은 튜닝 수치만 내세우지 않는다 ─────────
def test_accuracy_line_discloses_holdout_next_to_tuning_recall() -> None:
    replays = load_replays()
    assert service.holdout_recall(replays) == (4, 26)
    n1 = service.accuracy_line(replays["N1"], replays)
    assert "87.5%" in n1 and "(7/8" in n1 and "튜닝 노드" in n1 and "15.4%" in n1 and "(4/26)" in n1
    n2 = service.accuracy_line(replays["N2"], replays)
    assert "22.2%" in n2 and "홀드아웃 노드" in n2 and "15.4%" not in n2
    assert "측정하지 않았습니다" in service.accuracy_line(replays["P1"], replays)


def test_system_note_reads_model_and_cost_from_result() -> None:
    replays = load_replays()
    note = service.system_note(replays["N1"])
    cost = f"${replays['N1'].meta['cost_usd']:.3f}"  # O-2 재캡처로 $0.787 → $0.812 — 하드코딩하지 않는다
    assert replays["N1"].meta["model_id"] in note and cost in note and "에스코어 드림" in note
    gold = service.Result(meta={"source": "gold", "model_id": None})
    assert service.system_note(gold) == "화면에는 S-Core에서 제공한 에스코어 드림 폰트가 적용되어 있습니다."


# ── verifier 시연 토글 (지시문 O-3) ───────────────────────────────────────────
def _review_count_in_xlsx(data: bytes) -> int:
    import io

    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(data))["신뢰도"]
    return sum(ws.cell(row=r, column=2).value == "review" for r in range(2, ws.max_row + 1))


def test_demo_injection_flags_one_row_and_leaves_original_untouched() -> None:
    original = load_replays()["P1"]
    records_before = original.records
    first_id = id(original.records[0])
    rec0_before = list(original.records[0].recommendations)

    demo = service.demo_injected(original)
    table = service.worksheet_table(demo)
    red = [row for row in table if row[service.CONFIDENCE_COLUMN].startswith("🔴")]
    assert len(red) == 1 and red[0] is table[0]
    assert "unverified_standard: KOSHA GUIDE P-999" in red[0][service.FLAG_COLUMN]
    assert "review 1건 (규격 1·수치 0)" in service.summary_line(demo)

    # 원본은 그대로: 같은 리스트·같은 레코드 객체·같은 권고, 표에 🔴 0
    assert original.records is records_before and id(original.records[0]) == first_id
    assert original.records[0].recommendations == rec0_before
    assert not any(r[service.CONFIDENCE_COLUMN].startswith("🔴") for r in service.worksheet_table(original))
    assert "review 0건" in service.summary_line(original)
    # 다운로드(원본으로 만든다)의 신뢰도 시트엔 review 0
    assert _review_count_in_xlsx(service.export_files(original)["xlsx"][1]) == 0


def test_demo_injection_is_noop_for_gold() -> None:
    """골드 재생은 사람 작성 레코드라 검증 대상이 아니다 — 삽입하지 않는다(화면도 `is_gold` 면 토글을 안 그린다)."""
    gold = service.Result(meta={"source": "gold"}, records=load_replays()["P1"].records)
    assert service.demo_injected(gold) is gold


def test_app_verifier_demo_toggle(monkeypatch: pytest.MonkeyPatch) -> None:
    """토글 on → 🔴 1행·배너·요약 1건, 다운로드 xlsx 는 review 0 / off → 원래대로. 골드 공정엔 토글 없음."""
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    exported: list[str] = []  # 다운로드를 만든 Result 의 레코드 권고 전문 — 시연 삽입이 섞이면 안 된다
    real_export = service.export_files

    def spy(result: service.Result, edits: Any = None) -> dict[str, tuple[str, bytes]]:
        exported.append(" ".join(t for r in result.records for t in r.recommendations))
        return real_export(result, edits)

    monkeypatch.setattr(service, "export_files", spy)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    at.radio(key="mode").set_value(_CASES).run()
    at.selectbox(key="process_name").select("LPG 저장탱크 출하").run()
    boxes = [c for c in at.checkbox if c.label == service.DEMO_TOGGLE_LABEL]
    assert len(boxes) == 1 and boxes[0].value is False

    def red_rows() -> int:
        return sum(str(v).startswith("🔴") for v in at.dataframe[0].value[service.CONFIDENCE_COLUMN])

    def banner() -> bool:
        return any(service.DEMO_BANNER in m.value for m in at.markdown)

    assert red_rows() == 0 and not banner()
    boxes[0].check().run()
    assert not at.exception
    assert red_rows() == 1 and banner()
    assert any("review 1건 (규격 1·수치 0)" in c.value for c in at.caption)
    assert exported and "P-999" not in exported[-1]  # 토글 on 인 실행에서도 다운로드는 원본
    boxes = [c for c in at.checkbox if c.label == service.DEMO_TOGGLE_LABEL]
    boxes[0].uncheck().run()
    assert red_rows() == 0 and not banner()


# ── UX v3 (PRD_UX_v3 §4-4) — 문장 모드가 기본이고, 예시 칩이 입력 칸을 채운다 ─────────────
def test_app_starts_in_sentence_mode_with_empty_input(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    assert at.radio(key="mode").value == "문장으로 새 공정 분석"
    assert at.session_state["quick_text"] == ""
    assert len(at.dataframe) == 0
    assert not any("빠른 실호출" in b.label or "직접 입력" in b.label for b in at.button)


def test_app_example_chip_fills_input(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    chips = [b for b in at.button if (b.key or "").startswith("chip_")]
    assert len(chips) == 3
    at.button(key="chip_0").click().run()
    assert not at.exception
    assert at.session_state["quick_text"].startswith("수소충전소에서 튜브트레일러")
    assert at.text_area[0].value == at.session_state["quick_text"]
    assert next(b for b in at.button if b.label.startswith(_GENERATE)).disabled is False


# ── R-11 진행 알림 (본선 T-11·T-12) ───────────────────────────────────────────
def test_run_quick_reports_each_stage_in_order() -> None:
    seen: list[tuple[str, dict[str, Any]]] = []
    result = service.run_quick(_SENTENCE, "More", load_replays()["N1"], _MOCK_ENV, lambda e, p: seen.append((e, p)))
    assert [e for e, _ in seen] == ["parsed", "parameters", "guideword"]
    assert seen[0][1]["node_meta"].substance == "수소"
    assert seen[1][1]["parameters"] == result.meta["parameters"]
    assert len(seen[2][1]["records"]) == len(result.records)


def test_run_quick_survives_progress_callback_errors() -> None:
    def boom(*_: object) -> None:
        raise RuntimeError("화면 오류")

    result = service.run_quick(_SENTENCE, "More", load_replays()["N1"], _MOCK_ENV, boom)
    assert result.records  # AC-11-3 — 해석 단계 알림 예외도 생성을 멈추지 않는다


def test_service_uses_only_public_generator_api() -> None:
    source = (REPLAY_DIR.parents[1] / "apps" / "web" / "service.py").read_text(encoding="utf-8")
    assert "generator._" not in source  # R-11 AC-11-5

def test_progress_text_reports_stage_outputs() -> None:
    meta = NodeMeta(node="X1", substance="수소", phase="gas", equipment=["압축기", "디스펜서"])
    assert service.progress_text("parsed", {"node_meta": meta, "parsed_by": "llm"}, "More") == (
        0, "✅ 1/3 입력 해석 — 수소 · gas · 설비 압축기, 디스펜서", service.STAGE_PENDING[1],
    )
    step, text, pending = service.progress_text("parameters", {"parameters": [f"p{i}" for i in range(8)]}, "More")
    assert (step, text) == (1, "✅ 2/3 파라미터 8개 — p0, p1, p2, p3, p4, p5 외 2개") and "'More'" in str(pending)
    assert service.progress_text("guideword", {"records": [1, 2], "done": 3, "total": 7, "guideword": "Less"}, None)[1] == (
        "⏳ 3/3 가이드워드 3/7 판정 완료 — 방금 Less · 이탈 2건"
    )


def test_app_writes_stages_as_they_finish(monkeypatch: pytest.MonkeyPatch) -> None:
    """실패 경로는 상태 상자를 펼친 채 남긴다 — 실패 직전까지 끝난 단계가 화면에 남아 있어야 한다(T-12)."""
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)

    def partial_then_fail(text: str, guideword: str, replay: object, *_: object, on_progress: Any = None) -> None:
        on_progress("parsed", {"node_meta": NodeMeta(substance="수소", phase="gas", equipment=["압축기"]),
                               "parsed_by": "llm"})
        on_progress("parameters", {"parameters": ["유량", "압력"]})
        raise TimeoutError("판정 중 끊김")

    monkeypatch.setattr(service, "run_quick", partial_then_fail)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    at.text_area[0].input(_SENTENCE).run()
    at.button(key="run_quick").click().run()
    shown = [m.value for m in at.markdown]
    assert "✅ 1/3 입력 해석 — 수소 · gas · 설비 압축기" in shown
    assert "✅ 2/3 파라미터 2개 — 유량, 압력" in shown
    assert "⏳ 3/3 가이드워드 'More' 판정 중 (약 40~50초)" in shown
    assert not any(v.startswith("2/3 점검") for v in shown)  # 끝난 뒤 한꺼번에 찍던 옛 문구는 없다

# ── 지시문 V 부스 모드 (본선 F-04) ─────────────────────────────────────────────
def _booth_app(monkeypatch: pytest.MonkeyPatch, booth: bool, session_runs: int = 0) -> Any:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30)
    if booth:
        at.query_params["booth"] = "1"
    at.session_state["live_runs"] = session_runs
    return at.run()


def _chips(at: Any) -> list[str]:
    return [b.label for b in at.button if b.key and b.key.startswith("chip_")]


def _cta(at: Any) -> Any:
    return next(b for b in at.button if b.label.startswith(_GENERATE))


def test_default_screen_keeps_original_chips(monkeypatch: pytest.MonkeyPatch) -> None:
    assert _chips(_booth_app(monkeypatch, booth=False)) == ["수소충전소", "메탄올 하역", "실란 가스 캐비닛"]  # V-1


def test_booth_swaps_chips_and_fills_input(monkeypatch: pytest.MonkeyPatch) -> None:
    at = _booth_app(monkeypatch, booth=True)
    assert _chips(at) == ["학교 실험실 수소", "아파트 LPG 공급", "수영장 염소 소독"]  # V-2
    next(b for b in at.button if b.label == "수영장 염소 소독").click().run()
    assert "차아염소산나트륨" in at.text_area[0].value
    assert any(c.value.startswith("부스 모드 · 오늘 남은") for c in at.caption)  # V-5


def test_booth_ignores_session_limit_but_keeps_daily_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    at = _booth_app(monkeypatch, booth=True, session_runs=service.SESSION_LIMIT)
    next(b for b in at.button if b.label == "아파트 LPG 공급").click().run()
    assert not _cta(at).disabled  # V-3 세션 상한 미적용
    plain = _booth_app(monkeypatch, booth=False, session_runs=service.SESSION_LIMIT)
    next(b for b in plain.button if b.label == "수소충전소").click().run()
    assert _cta(plain).disabled  # 기본 화면은 세션 상한 그대로
    service._daily_runs[date.today()] = service.DAILY_LIMIT
    at = _booth_app(monkeypatch, booth=True)
    next(b for b in at.button if b.label == "아파트 LPG 공급").click().run()
    assert _cta(at).disabled  # 일 상한 = 비용 상한은 부스에서도 건다

def test_booth_pool_chip_selects_as_well_as(monkeypatch: pytest.MonkeyPatch) -> None:
    """V-6: 수영장 칩은 As well as 를 고른다(More 로는 산 혼입→염소가스 행이 안 나온다, 10/8 실측)."""
    at = _booth_app(monkeypatch, booth=True)
    assert at.selectbox[0].value == "More"
    next(b for b in at.button if b.label == "수영장 염소 소독").click().run()
    assert at.selectbox[0].value == "As well as"
    at.selectbox[0].select("Reverse").run()
    next(b for b in at.button if b.label == "아파트 LPG 공급").click().run()
    assert at.selectbox[0].value == "Reverse"  # 가이드워드를 지정하지 않은 칩은 사용자의 선택을 건드리지 않는다

# ── export-formats R-10 검토 반영 (본선 T-07·T-08) ─────────────────────────────
def _quick_result() -> service.Result:
    return service.run_quick(_SENTENCE, "More", load_replays()["N1"], _MOCK_ENV)


def _sheets(files: dict[str, tuple[str, bytes]]) -> Any:
    import io

    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(files["xlsx"][1]))


def test_export_without_review_is_unchanged() -> None:
    result = _quick_result()
    plain, empty = service.export_files(result), service.export_files(result, {})
    assert _sheets(plain).sheetnames == _sheets(empty).sheetnames == list(SHEET_ORDER)  # AC-10-1
    only_pending = service.export_files(result, {0: {service.REVIEW_COLUMN: "미검토"}})
    assert _sheets(only_pending).sheetnames == list(SHEET_ORDER)


def test_reject_removes_row_everywhere_and_logs_it() -> None:
    result = _quick_result()
    table = service.worksheet_table(result)
    top = max(range(len(table)), key=lambda i: table[i]["위험도"])  # LOPA 1순위 후보를 기각
    kept, log = service.apply_review(result, {top: {service.REVIEW_COLUMN: "기각"}})
    assert len(kept) == len(table) - 1 and log == [
        (table[top]["No"], "기각", table[top]["가이드워드"], table[top]["이탈"], "")
    ]
    wb = _sheets(service.export_files(result, {top: {service.REVIEW_COLUMN: "기각"}}))
    main = [r for r in wb["HAZOP워크시트"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]
    assert len(main) == len(table) - 1 and table[top]["이탈"] not in [r[3] for r in main]  # AC-10-4
    review = list(wb[REVIEW_SHEET].iter_rows(values_only=True))
    assert review[0] == REVIEW_HEADERS and review[1][1] == "기각"


def test_edit_changes_values_and_risk_formula_inputs() -> None:
    result = _quick_result()
    table = service.worksheet_table(result)
    original_s = table[0]["S(1-5)"]
    new_s = 1 if original_s != 1 else 2
    edits = {0: {"S(1-5)": new_s, "권고": "인터록 추가 · 절차서 개정", "결과": table[0]["결과"]}}  # 결과는 같은 값
    kept, log = service.apply_review(result, edits)
    assert kept[0]["S"] == new_s and kept[0]["recommendations"] == ["인터록 추가", "절차서 개정"]
    assert log == [(1, "수정", table[0]["가이드워드"], table[0]["이탈"], "권고, S(1-5)")]  # 같은 값은 수정 아님
    first = next(_sheets(service.export_files(result, edits))["HAZOP워크시트"].iter_rows(min_row=2, max_row=2))
    # AC-10-3 위험도는 수식이 다시 계산 — Y-2: 직접 입력은 C-C-37 대조표라 곱이 아니라 INDEX 조회
    assert first[7].value == new_s and first[9].value == "=INDEX({1,1,2;1,2,3;2,4,4;3,5,5},H2,I2)"


@pytest.mark.parametrize("bad", [0, 4, 2.5])
def test_edit_rejects_out_of_range_grade(bad: float) -> None:
    """Y-2: C-C-37 의 F 는 1~3 — 4 는 NH3 기준에선 맞아도 여기선 거부."""
    with pytest.raises(ValueError, match="1~3"):
        service.apply_review(_quick_result(), {0: {"F(1-5)": bad}})


def test_edit_range_follows_gold_criteria_for_gold_nodes() -> None:
    n1 = load_replays()["N1"]
    service.apply_review(n1, {0: {"F(1-5)": 5}})  # NH3 기준 1~5 — 통과
    with pytest.raises(ValueError, match="1~5"):
        service.apply_review(n1, {0: {"F(1-5)": 6}})


def test_app_passes_review_edits_to_downloads(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    seen: list[Any] = []
    real_export = service.export_files

    def spy(result: service.Result, edits: Any = None) -> dict[str, tuple[str, bytes]]:
        seen.append(edits)
        return real_export(result, edits)

    monkeypatch.setattr(service, "export_files", spy)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    at.text_area[0].input(_SENTENCE).run()
    next(b for b in at.button if b.label.startswith(_GENERATE)).click().run()
    assert at.dataframe and list(at.dataframe[0].value.columns)[0] == service.REVIEW_COLUMN  # 검토 열이 맨 앞
    result = at.session_state["quick_result"]
    key = service.review_key(result)
    at.session_state[key] = {"edited_rows": {0: {service.REVIEW_COLUMN: "기각"}}, "added_rows": [], "deleted_rows": []}
    at.run()
    assert seen[-1] == {0: {service.REVIEW_COLUMN: "기각"}}
    assert any("기각 1건" in c.value for c in at.caption)

# ── 지시문 W 외부 공개 HAZOP 대조 (LPG 충전소, IOCL 2014) ─────────────────────
def test_external_reference_is_pairs_only_and_split_external() -> None:
    from tools import capture_replay

    ref = json.loads((REPLAY_DIR.parents[1] / "data" / "reference" / "iocl_lpg_2014.json").read_text(encoding="utf-8"))
    assert {tuple(sorted(r)) for r in ref["records"]} == {("guideword", "node", "parameter", "source_node")}  # 문장 없음
    for node in ("I1", "I2", "I3", "I4"):
        assert capture_replay.default_split(node) == "external"
        assert len(capture_replay._gold_rows(node, "external")) == 11
    with pytest.raises(ValueError, match="골드셋이 없는"):
        capture_replay.capture_gold("I1", "external")  # 대조 기준은 레코드가 아니다 — 골드 재생 불가


def test_external_result_is_labelled_and_kept_out_of_nh3_table() -> None:
    replays = load_replays()
    meta = {**replays["N2"].meta, "node": "I2", "split": "external",
            "recall": {"recall": 9 / 11, "matched": 9, "total": 11}}
    result = service.Result(meta=meta, records=replays["N2"].records)
    assert "외부 공개 HAZOP 대비 recall 0.818 (9/11)" in service.summary_line(result)
    line = service.accuracy_line(result, replays)
    assert "외부 공개 HAZOP 대비" in line and "후하게" in line  # 쉬운 기준이라는 고지가 함께 붙는다
    assert service.criteria_notice(result) == service.CRITERIA_NOTICE
    table = service.evaluation_table({**replays, "I2": result})
    assert not any("I2" in str(row["노드"]) for row in table)  # NH3 골드셋 표와 섞지 않는다
    assert service.holdout_recall({**replays, "I2": result}) == service.holdout_recall(replays)


def test_app_shows_external_badge_instead_of_no_gold(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    at.radio(key="mode").set_value(_CASES).run()
    at.selectbox(key="process_name").select("LPG 충전소 (공개 HAZOP 대조)").run()
    shown = " ".join(m.value for m in at.markdown)
    assert "공개 HAZOP 대조 · 외부 팀 작성" in shown and "골드셋 없음" not in shown
    labels = [m.label for m in at.metric]
    assert "외부 대조 recall" in labels and "전문가 대비 recall" not in labels
    at.selectbox(key="process_name").select("LPG 저장탱크 출하").run()
    shown = " ".join(m.value for m in at.markdown)
    assert "예시 공정 · 정성 검토용" in shown and "골드셋 없음" not in shown


# ── 지시문 X 전체 가이드워드 · 보기 전환 · 공정 전체 · 검토 표 가독성 ─────────────────
def _view(at: Any) -> Any:
    """결과 보기 radio — 키가 view_<n> 으로 바뀐다(칩이 보기를 정하면 새 위젯)."""
    return next(r for r in at.radio if (r.key or "").startswith("view_"))


def _full_result() -> service.Result:
    return service.run_live(_SENTENCE, load_replays()["N1"], _MOCK_ENV)


def test_run_live_streams_partial_records_in_process_order() -> None:
    """X-1c·Y-1: 가이드워드 판정이 끝날 때마다 부분 레코드가 오고, 늘기만 하며, 공정 순서
    (파라미터 열거 순서 → 가이드워드 축 순서)다. 최종 결과도 같은 순서."""
    from core.agent.generate import STANDARD_GUIDEWORDS

    seen: list[dict[str, Any]] = []

    def spy(event: str, payload: dict[str, Any]) -> None:
        if event == "guideword":
            seen.append(payload)

    result = service.run_live(_SENTENCE, load_replays()["N1"], _MOCK_ENV, spy)
    assert [p["done"] for p in seen] == list(range(1, 8)) and all(p["total"] == 7 for p in seen)
    sizes = [len(p["records"]) for p in seen]
    assert sizes == sorted(sizes) and sizes[-1] == len(result.records)
    params = result.meta["parameters"]
    for records in (seen[-1]["records"], result.records):
        order = [(params.index(r.parameter), STANDARD_GUIDEWORDS.index(r.guideword)) for r in records]
        assert order == sorted(order)
    assert [r.id for r in result.records] == [f"x1-{i:03d}" for i in range(1, len(result.records) + 1)]
    m = result.meta
    assert m["guidewords"] == STANDARD_GUIDEWORDS and len(m["raw_calls"]) == 9 and m["parameters"]
    assert service.cells_label(result) == f"{m['judged_cells']}/{m['expected_cells']}"


def test_app_draws_partial_table_per_guideword(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    calls: list[int] = []
    real = service.partial_rows

    def spy(records: list[Any]) -> list[dict[str, object]]:
        calls.append(len(records))
        return real(records)

    monkeypatch.setattr(service, "partial_rows", spy)
    at = AppTest.from_file(str(_APP), default_timeout=60).run()
    at.text_area[0].input(_SENTENCE).run()
    _cta(at).click().run()
    assert not at.exception and len(calls) == 7 and calls == sorted(calls)


def test_failed_guideword_line() -> None:
    result = _full_result()
    assert service.failed_guidewords_line(result) is None
    broken = service.Result(meta={**result.meta, "review_guidewords": ["Reverse"]}, records=result.records)
    line = service.failed_guidewords_line(broken)
    assert line is not None and line.startswith("가이드워드 Reverse 판정 실패(review) — 나머지 6종은 정상")
    quick = service.Result(meta={**broken.meta, "source": "quick"}, records=result.records)
    assert service.failed_guidewords_line(quick) is None


@pytest.mark.parametrize("view", service.VIEWS)
def test_views_only_reorder_rows(view: str) -> None:
    """X-2b: 어느 보기든 같은 행 집합을 정확히 한 번씩 가리킨다."""
    result = _full_result()
    groups = service.view_groups(result, view)
    indices = [i for _, _, idx in groups for i in idx]
    assert sorted(indices) == list(range(len(result.records)))
    records = service.verified(result)[0]
    if view == service.VIEWS[1]:
        assert [g[0] for g in groups] == result.meta["guidewords"]
        assert all(records[i].guideword == name for name, _, idx in groups for i in idx)
        n_params = len(result.meta["parameters"])
        for name, head, idx in groups:
            skipped = n_params - len({records[i].parameter for i in idx})
            assert head == f"{name} — {len(idx)}건" + (f" (해당 없음 {skipped}셀)" if skipped else "")
    if view == service.VIEWS[2]:
        params = [p for p in result.meta["parameters"] if any(r.parameter == p for r in records)]
        assert [g[0] for g in groups] == params
        assert all(records[i].parameter == name for name, _, idx in groups for i in idx)


def test_app_view_switch_keeps_edits_in_worksheet_view_only(monkeypatch: pytest.MonkeyPatch) -> None:
    """X-2b: 편집은 워크시트 순서에서만. 가이드워드별은 읽기 전용 묶음이고 검토 결과가 같은 행에 붙는다."""
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=60).run()
    at.text_area[0].input(_SENTENCE).run()
    _cta(at).click().run()
    result = at.session_state["quick_result"]
    table = service.worksheet_table(result)
    target = len(table) - 1  # 마지막 행을 기각 — 묶음 보기에서 같은 No 에 붙어야 한다
    key = service.review_key(result)
    at.session_state[key] = {"edited_rows": {target: {service.REVIEW_COLUMN: "기각"}}, "added_rows": [], "deleted_rows": []}
    at.run()
    _view(at).set_value(service.VIEWS[1]).run()
    assert not at.exception
    frames = at.dataframe
    assert len(frames) == len([g for g in service.view_groups(result, service.VIEWS[1]) if g[2]])
    rejected = [
        f.value["No"].iloc[j] for f in frames for j, v in enumerate(f.value[service.REVIEW_COLUMN]) if v == "기각"
    ]
    assert rejected == [table[target]["No"]]
    assert any("워크시트 순서" in c.value for c in at.caption)
    assert any("기각 1건" in c.value for c in at.caption)  # 다운로드 반영은 보기와 무관
    # 돌아오면 편집 표가 저장본으로 다시 그려진다 — 보기를 오가도 검토가 사라지지 않는다.
    # 그려지지 않은 편집 표의 위젯 상태는 Streamlit 이 지운다 — 그래도 저장본으로 반영된다.
    at.run()
    assert key not in at.session_state
    assert any("기각 1건" in c.value for c in at.caption)
    _view(at).set_value(service.VIEWS[0]).run()
    assert not at.exception
    assert list(at.dataframe[0].value[service.REVIEW_COLUMN]).count("기각") == 1
    assert at.dataframe[0].value[service.REVIEW_COLUMN].iloc[target] == "기각"
    assert any("기각 1건" in c.value for c in at.caption)


def test_merge_and_with_edits() -> None:
    merged = service.merge_edits({0: {"검토": "기각"}}, {"0": {"S(1-5)": 2}, "3": {"검토": "채택"}})
    assert merged == {0: {"검토": "기각", "S(1-5)": 2}, 3: {"검토": "채택"}}
    table = [{"검토": "미검토", "S(1-5)": 4}, {"검토": "미검토", "S(1-5)": 5}]
    assert service.with_edits(table, {1: {"검토": "기각", "없는 열": 1}}) == [
        {"검토": "미검토", "S(1-5)": 4}, {"검토": "기각", "S(1-5)": 5},
    ]
    assert table[1]["검토"] == "미검토"  # 원본은 그대로


def test_combine_replays_keeps_node_order_and_split_recalls() -> None:
    replays = load_replays()
    nh3 = next(p for p in service.CATALOG if p["id"] == "nh3_sts")
    combined = service.combine_replays(nh3, replays)
    assert combined is not None and combined.meta["nodes"] == ["N1", "N2", "N3", "N4"]
    assert [r.id for r in combined.records] == [r.id for n in ("N1", "N2", "N3", "N4") for r in replays[n].records]
    assert len({r.id for r in combined.records}) == len(combined.records)  # 노드 접두사로 id 충돌 없음
    assert combined.meta["recall_parts"] == [("튜닝", 7, 8), ("홀드아웃", 4, 26)]  # 섞지 않는다
    assert service.recall_tiles(combined) == [("튜닝 recall", "0.875"), ("홀드아웃 recall", "0.154")]
    assert service.criteria_notice(combined) is None and not combined.is_gold
    lpg = service.combine_replays(next(p for p in service.CATALOG if p["id"] == "lpg_iocl"), replays)
    assert lpg is not None and service.recall_tiles(lpg) == [("외부 대조 recall", "0.909")]  # 40/44
    assert "후한 기준" in service.accuracy_line(lpg, replays)
    assert service.export_files(combined)["xlsx"][1][:2] == b"PK"  # 다운로드 1파일


def test_app_whole_process_replay(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=60).run()
    at.radio(key="mode").set_value(_CASES).run()
    at.button(key="preset_all").click().run()
    assert not at.exception
    nodes = {str(v).split(" ")[0] for v in at.dataframe[0].value["노드"]}
    assert nodes == {"N1", "N2", "N3", "N4"}  # 노드가 여러 개라 노드 열을 숨기지 않는다
    assert {"튜닝 recall", "홀드아웃 recall"} <= {m.label for m in at.metric}
    assert any(e.label == "노드별 출처" for e in at.expander)


def test_booth_pool_chip_opens_as_well_as_group(monkeypatch: pytest.MonkeyPatch) -> None:
    """X-4: 수영장 칩 → 가이드워드별 보기 · As well as 묶음만 펼침 · 보조 실행 기본값 As well as."""
    at = _booth_app(monkeypatch, booth=True)
    next(b for b in at.button if b.label == "수영장 염소 소독").click().run()
    assert at.session_state["view_pref"] == service.VIEWS[1]
    assert at.selectbox(key="guideword").value == "As well as"
    _cta(at).click().run()
    assert not at.exception
    assert _view(at).value == service.VIEWS[1] and _view(at).index == 1  # 화면 선택 표시도 같은 값(index 로 생성)
    groups = [e for e in at.expander if e.label.split(" — ")[0] in service.GUIDEWORD_DEFINITIONS]
    assert len(groups) == 7
    assert [e.label.split(" — ")[0] for e in groups if e.proto.expanded] == ["As well as"]


# ── 검토 표 가독성 (10/8 사용자 피드백 — 검토자는 열을 눈으로 따라가며 판단한다) ─────────
def test_review_columns_put_decision_columns_first_and_hide_uniform() -> None:
    result = _full_result()
    table = [{service.REVIEW_COLUMN: "미검토", **row} for row in service.readable_rows(service.worksheet_table(result))]
    order = service.review_column_order(table)
    assert order[:8] == [service.REVIEW_COLUMN, "No", "가이드워드", "이탈", "위험도", "S(1-5)", "F(1-5)", "신뢰도"]
    assert "노드" not in order and "시나리오 연계" not in order and service.FLAG_COLUMN not in order
    assert {"이탈", "원인", "결과", "권고", "S(1-5)", "F(1-5)", "위험도"} <= set(order)
    demo = service.readable_rows(service.worksheet_table(service.demo_injected(load_replays()["P1"])))
    assert service.FLAG_COLUMN in service.review_column_order(demo)  # 플래그가 있으면 보인다
    listed = next(r for r in table if "·" in str(r["원인"]))
    assert " · " in str(listed["원인"])


def test_spacing_only_is_not_an_edit() -> None:
    """화면은 ' · ' 로 띄워 보여 준다 — 그 값을 그대로 돌려받아도 '수정' 으로 치지 않는다."""
    result = _full_result()
    shown = service.readable_rows(service.worksheet_table(result))
    row = next(i for i, r in enumerate(shown) if " · " in str(r["원인"]))
    kept, log = service.apply_review(result, {row: {"원인": shown[row]["원인"], service.REVIEW_COLUMN: "채택"}})
    assert log == [(shown[row]["No"], "채택", shown[row]["가이드워드"], shown[row]["이탈"], "")]
    assert kept[row]["causes"] == service._display_records(result)[row]["causes"]

# ── 'HAZOP 이 처음이라면' 안내 (10/8 사용자 요청) ─────────────────────────────────
def test_guide_reuses_existing_definitions() -> None:
    """뜻·등급은 새로 쓰지 않고 생성 프롬프트·골드셋 평가기준에서 읽는다."""
    from apps.web import guide
    from core.agent.generate import GUIDEWORD_DEFINITIONS, STANDARD_GUIDEWORDS

    rows = guide.guideword_rows()
    assert [r["가이드워드"] for r in rows] == STANDARD_GUIDEWORDS
    assert all(r["뜻"] == GUIDEWORD_DEFINITIONS[r["가이드워드"]] for r in rows)
    from core.criteria import load_criteria

    gold, cc37 = load_criteria("nh3_sts_bunkering"), load_criteria("kosha_cc37_2026")
    assert [r["등급"] for r in guide.criteria_rating_rows(gold)] == [5, 4, 3, 2, 1]
    assert [b["판정"] for b in guide.criteria_band_rows(gold)] == ["높음", "중간(ALARP)", "낮음"]
    rows = guide.criteria_rating_rows(cc37)  # S 4단계·F 3단계 — 4등급의 F 칸은 비어 있다
    assert [r["등급"] for r in rows] == [4, 3, 2, 1] and rows[0]["빈도 — 얼마나 자주 일어날 수 있나"] == ""
    assert guide.criteria_matrix_rows(cc37)[0] == {"심각도＼빈도": "심각도 4", "빈도 1": 3, "빈도 2": 5, "빈도 3": 5}


def test_worked_example_risk_matches_band() -> None:
    """설명용 예시의 위험도·판정이 C-C-37 대조표·구간과 맞는다(예시만 고치고 표를 안 보는 실수 방지)."""
    import re

    from apps.web import guide
    from core.criteria import load_criteria

    value = {s: v for s, _, v in guide.WORKED_EXAMPLE}["심각도 · 빈도 → 위험도"]
    s, f, risk, verdict = re.match(r"심각도 (\d)\S* · 빈도 (\d)\S* → 위험도 (\d+) → (.+)", value).groups()  # type: ignore[union-attr]
    cc37 = load_criteria("kosha_cc37_2026")
    assert cc37.risk(int(s), int(f)) == int(risk) and cc37.band(int(risk)) == verdict


def test_app_renders_intro_tabs(monkeypatch: pytest.MonkeyPatch) -> None:
    from streamlit.testing.v1 import AppTest

    for key, value in _MOCK_ENV.items():
        monkeypatch.setenv(key, value)
    at = AppTest.from_file(str(_APP), default_timeout=30).run()
    assert not at.exception
    labels = [t.label for t in at.tabs]
    assert {"① 한 줄이 만들어지는 과정", "② 가이드워드 7종", "③ 위험도 (심각도·빈도)", "④ 결과 화면 읽는 법", "⑤ 회의와 비교"} <= set(labels)
    assert len(at.table) >= 3 and len(at.dataframe) == 0  # 안내 표는 st.table — 결과 표 시험과 섞이지 않는다


# ── 지시문 Y-1 No 공정 순서 ───────────────────────────────────────────────────
def _process_key_ok(records: list[Any], params: list[str]) -> bool:
    from core.agent.generate import PROCEDURAL_GUIDEWORDS, STANDARD_GUIDEWORDS

    axis = STANDARD_GUIDEWORDS + PROCEDURAL_GUIDEWORDS
    names = list(dict.fromkeys([*params, *(r.parameter for r in records)]))
    keys = [(names.index(r.parameter), axis.index(r.guideword)) for r in records]
    return keys == sorted(keys)


def test_replays_load_in_process_order_and_gold_untouched() -> None:
    """Y-1: 생성 재생은 불러올 때 파라미터 → 가이드워드 순. 파일 id 는 그대로(대조용), 골드는 원래 순서."""
    on_disk = {p.stem: json.loads(p.read_text(encoding="utf-8")) for p in REPLAY_DIR.glob("*.json")}
    replays = load_replays()
    for node, result in replays.items():
        raw = next(v for v in on_disk.values() if v.get("node") == node and v.get("captured_at") == result.meta.get("captured_at"))
        if result.is_gold:
            assert [r.id for r in result.records] == [r["id"] for r in raw["records"]]
            continue
        assert _process_key_ok(result.records, result.meta.get("parameters") or [])
        assert sorted(r.id for r in result.records) == sorted(r["id"] for r in raw["records"])


def test_screen_xlsx_review_index_share_process_order() -> None:
    """Y-G1: 화면 '워크시트 순서'·Excel No·검토 편집 인덱스가 같은 공정 순서."""
    result = _full_result()
    assert _process_key_ok(result.records, result.meta["parameters"])
    table = service.worksheet_table(result)
    wb = _sheets(service.export_files(result))
    main = [r for r in wb["HAZOP워크시트"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]
    assert [r[0] for r in main] == [t["No"] for t in table] == list(range(1, len(table) + 1))
    assert [(r[2], r[3]) for r in main] == [(t["가이드워드"], t["이탈"]) for t in table]
    assert [t["이탈"] for t in table] == [r.deviation for r in result.records]
    # 검토 편집 인덱스 0 = 화면 첫 행 = xlsx 첫 행
    kept, _ = service.apply_review(result, {0: {service.REVIEW_COLUMN: "기각"}})
    assert table[0]["이탈"] not in [k["deviation"] for k in kept]
