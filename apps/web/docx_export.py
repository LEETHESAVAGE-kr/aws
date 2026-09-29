"""LOPA 초안 Markdown → Word(.docx) 변환 — 화면 다운로드 전용 (사용자 결정 9/29, spec export-formats 밖).

정본은 `core/export/lopa.py` 의 Markdown 이다. 이 모듈은 그 출력이 쓰는 문법(`#`~`###` 제목, `>` 인용,
`- ` 목록, `|` 표, 일반 문단, `**굵게**`, 인라인 코드 백틱)만 옮긴다. 마크업 기호 외 글자는 바꾸지 않는다.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_TABLE_RULE = re.compile(r"^\|[\s:|-]+\|$")
_FONT = "맑은 고딕"
_MUTED = RGBColor(0x6A, 0x6A, 0x6A)


def _add_runs(paragraph, text: str) -> None:  # noqa: ANN001 — docx Paragraph
    """`**굵게**` 는 굵게, 인라인 코드 백틱은 떼고, 나머지는 원문 그대로."""
    for i, part in enumerate(_BOLD.split(text.replace("`", ""))):
        if part:
            paragraph.add_run(part).bold = i % 2 == 1


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def lopa_markdown_to_docx(markdown: str) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", _FONT)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                if not _TABLE_RULE.match(lines[i].strip()):
                    block.append(_cells(lines[i]))
                i += 1
            table = doc.add_table(rows=len(block), cols=max(len(r) for r in block))
            table.style = "Table Grid"
            for r, row in enumerate(block):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], value)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        heading = re.match(r"^(#{1,3}) (.*)$", line)
        if heading:
            doc.add_heading(heading.group(2).replace("`", ""), level=len(heading.group(1)) - 1)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            _add_runs(p, line[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = _MUTED
        elif line.startswith("- "):
            _add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            _add_runs(doc.add_paragraph(), line)
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
